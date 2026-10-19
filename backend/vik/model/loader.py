"""Parse CSV/XLSX in any layout into column-profiled datasets.

Handles currency/sign/separator conventions (`$1,234`, `(5,000)`, `-$5`, `12%`, `1.234,56`), finds the
header row under any title rows, and detects the period axis in whatever form it comes: a date column
(ISO, US or day-first dates, month names, quarters, ISO weeks, fiscal years, Excel serials), separate
Year and Month columns, or periods across the header row (transposed). Several rows per period (daily
or weekly data, a ledger with one line per account) roll up to the period. Values a header or title says
are in thousands or millions are scaled. Tables keyed by period + customer or signup cohort become
dimension tables; a table with no dates stays raw.
"""

import calendar
import csv
import hashlib
import io
import math
import re
from collections import Counter
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Any

from vik.model.currencies import STRIP_RE, currency_in

# English month names from the stdlib (the C locale; nothing here calls setlocale): "january" … "december".
FULL_MONTHS = [calendar.month_name[i].lower() for i in range(1, 13)]
MONTHS = {m[:3]: i + 1 for i, m in enumerate(FULL_MONTHS)}
MONTHS.update({"sept": 9, FULL_MONTHS[5]: 6, FULL_MONTHS[6]: 7})  # the 4-letter spellings headers use

# A series row's source lines in the file ("row 6" in a figure's source); not a column.
SRC_ROWS = "_src_rows"

_NUM_RE = re.compile(r"^[+-]?\d+(\.\d+)?([eE][+-]?\d+)?$")
_MISSING = {"", "-", "—", "–", "n/a", "na", "nan", "null", "none", "#n/a", "#div/0!", "#value!", "#ref!"}
_SUFFIX = {"k": 1e3, "m": 1e6, "mm": 1e6, "b": 1e9, "bn": 1e9}


def parse_number(raw: Any) -> float | None:
    """Normalise a spreadsheet cell to a float; None when it isn't a number."""
    if raw is None:
        return None
    if isinstance(raw, bool):
        return None
    if isinstance(raw, int | float):
        return float(raw) if math.isfinite(raw) else None
    s = str(raw).strip()
    if s.lower() in _MISSING:
        return None
    neg, money = False, False
    s = s.replace("−", "-")
    # Peel signs, parentheses and currency markers in any order: "$(1,234.00)", "USD (1,234)", "-$5", "(US$5)".
    while s:
        if s in ("-", "–", "—"):  # Excel's "Accounting" zero: a lone dash beside a currency symbol
            return 0.0 if money else None
        bare = STRIP_RE.sub(
            "", s
        ).strip()  # US$, £, "USD", "1 200 €" … (the column's currency is read separately)
        if bare != s:
            money, s = True, bare
        elif s.startswith("(") and s.endswith(")"):
            neg, s = (not neg), s[1:-1].strip()
        elif s.startswith("-"):
            neg, s = (not neg), s[1:].strip()
        elif s.startswith("+"):
            s = s[1:].strip()
        elif s.endswith("-"):  # trailing minus, as some accounting exports print it
            neg, s = (not neg), s[:-1].strip()
        else:
            break
    pct = s.endswith("%")
    if pct:
        s = s[:-1].strip()
    m = re.match(r"^(.*?)\s*(mm|bn|[kmb])$", s, flags=re.I)
    core, suffix = (m.group(1), m.group(2).lower()) if m else (s, "")
    core = re.sub(r"[\s  ']", "", core)  # 1 234 / 1'234 thousands separators
    # Day-first locales: 1.234,56 · 1.234.567 · 12,5 (a comma followed by 1-2 digits is never a thousands separator)
    if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?|\d+,\d{1,2}", core) and ("," in core or core.count(".") > 1):
        core = core.replace(".", "").replace(",", ".")
    core = core.replace(",", "")
    if not _NUM_RE.match(core):
        return None
    v = float(core) * _SUFFIX.get(suffix, 1.0)
    if not math.isfinite(v):  # "1e400" overflows to infinity: not a figure anything can be computed from
        return None
    if pct:
        v /= 100.0
    return -v if neg else v


# ---------------------------------------------------------------- periods

_WEEKDAY = re.compile(r"^(mon|tue|wed|thu|fri|sat|sun)[a-z]*\.?,?\s+")
_DMY = re.compile(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2}|\d{4})(?:[ t].*)?$")
_ACTUAL = re.compile(r"\b(actuals?|act\.?)\b|\(a\)")


def _year(s: str) -> int:
    y = int(s)
    return y + 2000 if y < 100 else y


def _month_word(w: str) -> int | None:
    """January / jan / jan. / sept → 1-12; anything else (\"marketing\") → None."""
    w = w.strip(".").lower()
    if w in MONTHS:
        return MONTHS[w]
    if len(w) >= 3 and w[:3] in MONTHS and FULL_MONTHS[MONTHS[w[:3]] - 1].startswith(w):
        return MONTHS[w[:3]]
    return None


def _month(y: int, m: int) -> tuple[str, str] | None:
    return (f"{y:04d}-{m:02d}", "month") if 1 <= m <= 12 and 1900 <= y <= 2200 else None


def parse_period(raw: Any, dayfirst: bool | None = None) -> tuple[str, str] | None:
    """Return (normalised key, granularity) for a period label, else None.

    Months → "YYYY-MM" (a full date keeps its month; an ISO week, its Thursday's month); quarters →
    "YYYY-Qn"; years → "YYYY". `dayfirst` settles 03/04/2025: None reads it month-first unless a part
    is over 12 (columns decide it once for every cell, see `period_reader`).
    """
    if raw is None or isinstance(raw, bool):
        return None
    if hasattr(raw, "year") and hasattr(raw, "month"):
        return f"{raw.year:04d}-{raw.month:02d}", "month"
    s = str(raw).strip().lower().replace("_", " ")
    s = _ACTUAL.sub("", s)
    s = re.sub(r"\s+", " ", _WEEKDAY.sub("", s)).strip(" -,")
    if not s:
        return None
    m = re.match(r"^(\d{4})[-/.](\d{1,2})(?:[-/.](\d{1,2}))?(?:[ t].*)?$", s)
    if m:
        return _month(int(m.group(1)), int(m.group(2)))
    m = re.match(r"^(\d{4})[- ]?w(\d{1,2})$", s)
    if m:
        try:
            d = date.fromisocalendar(int(m.group(1)), int(m.group(2)), 4)
        except ValueError:
            return None
        return f"{d.year:04d}-{d.month:02d}", "month"
    m = _DMY.match(s)
    if m:
        a, b, y = int(m.group(1)), int(m.group(2)), _year(m.group(3))
        first = dayfirst if dayfirst is not None else a > 12
        day, mon = (a, b) if first else (b, a)
        return _month(y, mon) if 1 <= day <= 31 else None
    m = re.match(r"^(\d{1,2})[-/.](\d{4})$", s)
    if m:
        return _month(int(m.group(2)), int(m.group(1)))
    m = re.match(r"^([a-z]{3,9})\.?(?:[\s\-/]+(\d{1,2})(?:st|nd|rd|th)?)?[\s\-/',.]*(\d{2}|\d{4})$", s)
    if m and (mon := _month_word(m.group(1))):
        return _month(_year(m.group(3)), mon)
    m = re.match(r"^(\d{1,2})(?:st|nd|rd|th)?[\s\-/.]+([a-z]{3,9})\.?[\s\-/.,']+(\d{2}|\d{4})$", s)
    if m and (mon := _month_word(m.group(2))):
        return _month(_year(m.group(3)), mon)
    m = re.match(r"^(\d{4})[\s\-/.']+([a-z]{3,9})\.?$", s)
    if m and (mon := _month_word(m.group(2))):
        return _month(int(m.group(1)), mon)
    m = re.match(
        r"^(?:fy)?\s*(?:q|qtr|quarter)\s*([1-4])[\s\-'/,]*(?:fy)?\s*'?(\d{2}|\d{4})$", s
    ) or re.match(r"^([1-4])q[\s\-'/]*(?:fy)?\s*'?(\d{2}|\d{4})$", s)
    if m:
        return f"{_year(m.group(2)):04d}-Q{m.group(1)}", "quarter"
    m = re.match(r"^(?:fy)?\s*'?(\d{2}|\d{4})[\s\-'/]*(?:q|qtr|quarter)\s*([1-4])$", s)
    if m:
        return f"{_year(m.group(1)):04d}-Q{m.group(2)}", "quarter"
    m = re.match(r"^(?:fy|cy)\s*'?(\d{2}|\d{4})$", s) or re.match(r"^(\d{4})$", s)
    if m and 1990 <= _year(m.group(1)) <= 2100:
        return f"{_year(m.group(1)):04d}", "year"
    return None


def period_sort_key(p: str) -> tuple[int, int]:
    if "-Q" in p:
        y, q = p.split("-Q")
        return int(y), int(q) * 3
    if "-" in p:
        y, mo = p.split("-")
        return int(y), int(mo)
    return int(p), 12


def period_end_month(p: str) -> int:
    """The period's last month as a running month count (year × 12 + month index), for calendar arithmetic."""
    y, m = period_sort_key(p)
    return y * 12 + m - 1


def shift_period(p: str, months: int) -> str:
    """The period of the same granularity `months` calendar months away (2025-02, -12 → 2024-02; 2025-Q1, -3 → 2024-Q4)."""
    n = period_end_month(p) + months
    y, m = divmod(n, 12)
    if "-Q" in p:
        return f"{y}-Q{m // 3 + 1}"
    if "-" in p:
        return f"{y:04d}-{m + 1:02d}"
    return f"{y}"


def quarter_of(month_key: str) -> str:
    y, m = month_key.split("-")
    return f"{y}-Q{(int(m) - 1) // 3 + 1}"


def _blank(v: Any) -> bool:
    return v is None or str(v).strip() == ""


def _dayfirst(values: list[Any]) -> bool | None:
    """True when the column's numeric dates are day-first (a first part over 12), False when month-first.

    When every date is ambiguous (01/01/2025 … 01/12/2025), the part that varies is the month: a fixed
    first part with a moving second part is day-first, otherwise the dates would collapse into one month."""
    firsts: set[int] = set()
    seconds: set[int] = set()
    for v in values:
        m = _DMY.match(str(v).strip()) if isinstance(v, str) else None
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            if a > 12 >= b:
                return True
            if b > 12 >= a:
                return False
            firsts.add(a)
            seconds.add(b)
    if len(firsts) == 1 and len(seconds) > 1:
        return True
    if len(seconds) == 1 and len(firsts) > 1:
        return False
    return None


def _day_of(v: Any, dayfirst: bool | None) -> date | None:
    """The full date a period cell names (a datetime, 2025-01-31, 31/01/2025), else None (month labels)."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    try:
        m = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(?:[ t].*)?$", s, re.I)
        if m:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        m = _DMY.match(s.lower())
        if m:
            a, b, y = int(m.group(1)), int(m.group(2)), _year(m.group(3))
            first = dayfirst if dayfirst is not None else a > 12
            return date(y, b, a) if first else date(y, a, b)
    except ValueError:
        return None
    return None


# A header that is a date word ("Date", "Invoice date", "Billing month", "Period end"), so numbers under it may be dates.
_DATE_HEAD = re.compile(
    r"(?:[a-z]+[ _-])?(?:date|month|period|week|day|yyyymm|yyyy-?mm)(?:[ _-](?:end|ending|start|of))?", re.I
)


def _serial_period(v: Any) -> tuple[str, str] | None:
    """An Excel serial day number (45292 = 2024-01-01)."""
    n = parse_number(v)
    if n is None or not 20_000 <= n <= 80_000:
        return None
    d = date(1899, 12, 30) + timedelta(days=int(n))
    return f"{d.year:04d}-{d.month:02d}", "month"


def _compact_period(v: Any) -> tuple[str, str] | None:
    """202401 → 2024-01."""
    s = str(v).strip().split(".")[0]
    if not re.fullmatch(r"\d{6}", s):
        return None
    return _month(int(s[:4]), int(s[4:])) if 1990 <= int(s[:4]) <= 2100 else None


Reader = Callable[[Any], tuple[str, str] | None]


def period_reader(name: str, values: list[Any], threshold: float = 0.8) -> tuple[Reader, bool] | None:
    """(reader, by_header) when the column reads as dates. by_header: only its header says so (serials,
    YYYYMM numbers), so it would otherwise pass for a numeric column."""
    present = [v for v in values if not _blank(v)]
    if not present:
        return None
    first = _dayfirst(present)

    def read(v: Any) -> tuple[str, str] | None:
        return parse_period(v, first)

    if _share(present, lambda v: read(v) is not None) >= threshold:
        return read, False
    if _DATE_HEAD.fullmatch((name or "").strip()):
        for alt in (_compact_period, _serial_period):
            if _share(present, lambda v, f=alt: f(v) is not None) >= threshold:
                return alt, True
    return None


# ---------------------------------------------------------------- tables


@dataclass
class RawTable:
    name: str
    header: list[str]
    rows: list[list[Any]]
    preamble: str = ""  # title rows above the header ("All figures in CAD thousands")
    lines: list[int] = field(default_factory=list)  # each data row's 1-based line (row) in the file


@dataclass
class ColumnProfile:
    name: str
    dtype: str
    sample_values: list[str]
    null_rate: float
    currency: str | None = None  # the currency its header or cells name, if exactly one
    unit_hint: str | None = None  # "pct", "date" or "text" when the values say so
    scale: float = 1.0  # values were multiplied by this ("Revenue ($000)")

    def as_dict(self) -> dict:
        return {
            "name": self.name,
            "dtype": self.dtype,
            "sample_values": self.sample_values,
            "null_rate": round(self.null_rate, 3),
            "currency": self.currency,
            "unit_hint": self.unit_hint,
            "scale": self.scale,
        }


@dataclass
class ParsedDataset:
    name: str
    table_kind: str
    period_axis: str | None
    granularity: str
    transposed: bool
    columns: list[ColumnProfile]
    rows: list[dict]
    fingerprint: str
    notes: list[str] = field(default_factory=list)

    def column_names(self) -> list[str]:
        return [c.name for c in self.columns]


def read_tables(data: bytes, filename: str) -> list[RawTable]:
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(data, filename)
    return [_read_csv(data, filename)]


def _header_index(rows: list[list[Any]]) -> int:
    """The header row: the first row (of the top 30) about as wide as the table and mostly labels, so
    title rows ("Acme Inc. · Income statement", "In CAD thousands") above it are skipped."""
    counts = [sum(1 for c in r if not _blank(c)) for r in rows[:30]]
    if not counts:
        return 0
    widest = max(counts)
    for i, n in enumerate(counts):
        if n >= 2 and n >= 0.6 * widest:
            cells = [c for c in rows[i] if not _blank(c)]
            labels = sum(1 for c in cells if parse_number(c) is None or parse_period(c) is not None)
            if labels >= 0.5 * len(cells):
                return i
    return 0


def _split(name: str, rows: list[list[Any]], lines: list[int]) -> RawTable:
    if not rows:
        return RawTable(name, [], [])
    h = _header_index(rows)
    preamble = " · ".join(" ".join(str(c).strip() for c in r if not _blank(c)) for r in rows[:h])
    header = [str(c).strip() if not _blank(c) else "" for c in rows[h]]
    return RawTable(name, header, rows[h + 1 :], preamble, lines[h + 1 :])


def _read_csv(data: bytes, filename: str) -> RawTable:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows: list[list[Any]] = []
    lines: list[int] = []  # the line each kept row starts on, so sources cite the file's own line numbers
    last = 0
    for r in reader:
        start, last = last + 1, reader.line_num
        if any(c.strip() for c in r):
            rows.append(r)
            lines.append(start)
    if not rows:
        return RawTable(filename, [], [])
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    return _split(_stem(filename), rows, lines)


def _read_xlsx(data: bytes, filename: str) -> list[RawTable]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    tables: list[RawTable] = []
    for ws in wb.worksheets:
        numbered = [(n, list(r)) for n, r in enumerate(ws.iter_rows(values_only=True), start=1)]
        numbered = [(n, r) for n, r in numbered if any(c not in (None, "") for c in r)]
        if len(numbered) < 2:
            continue
        width = max(len(r) for _, r in numbered)
        rows = [r + [None] * (width - len(r)) for _, r in numbered]
        t = _split(f"{_stem(filename)} · {ws.title}", rows, [n for n, _ in numbered])
        t.header = [h or f"col_{i + 1}" for i, h in enumerate(t.header)]
        tables.append(t)
    return tables


def _stem(filename: str) -> str:
    base = re.split(r"[\\/]", filename)[-1]
    return base.rsplit(".", 1)[0] if "." in base else base


def fingerprint(columns: list[str]) -> str:
    norm = "|".join(sorted(re.sub(r"[^a-z0-9]", "", c.lower()) for c in columns))
    return hashlib.sha256(norm.encode()).hexdigest()[:16]


def _share(values: list[Any], pred) -> float:
    vals = [v for v in values if not _blank(v)]
    if not vals:
        return 0.0
    return sum(1 for v in vals if pred(v)) / len(vals)


# "($000)", "(in $ thousands)", "in USD millions", "(k$)"; a symbol has no word boundary before it, so only
# the three-letter code alternative takes `\b`.
_SCALE_RES = [
    (
        re.compile(
            r"\$\s?bn?\b|\bin billions\b|\(\s*(?:in\s+)?(?:[a-z]{3}\s*)?(?:[$£€]\s*)?(?:bn|b|billions?)\s*\)|"
            r"(?:\b[a-z]{3}|[$£€])\s?billions?\b|\bb(?:n)?\s?\$",
            re.I,
        ),
        1e9,
    ),
    (
        re.compile(
            r"\$\s?(?:mm|m)\b|\bin millions\b|\(\s*(?:in\s+)?(?:[a-z]{3}\s*)?(?:[$£€]\s*)?(?:mm|m|mn|millions?)\s*\)|"
            r"(?:\b[a-z]{3}|[$£€])\s?millions?\b|\bmm?\s?\$",
            re.I,
        ),
        1e6,
    ),
    (
        re.compile(
            r"\$\s?k\b|\$\s?'?000s?\b|'000s?\b|\bin thousands\b|"
            r"\(\s*(?:in\s+)?(?:[a-z]{3}\s*)?(?:[$£€]\s*)?(?:k|000s?|thousands?)\s*\)|"
            r"(?:\b[a-z]{3}|[$£€])\s?(?:'?000s?|thousands?)\b|\bk\s?\$",
            re.I,
        ),
        1e3,
    ),
]
_NOT_MONEY = re.compile(
    r"%|percent|margin|rate|ratio|headcount|head count|employees|\bfte|staff|customers|clients|"
    r"users|subscribers|accounts|logos|count|number of|#|months|days|units",
    re.I,
)


def scale_of(text: str) -> float:
    """1e3 / 1e6 / 1e9 when a header or title says the values are in thousands / millions / billions."""
    for rx, mult in _SCALE_RES:
        if rx.search(text or ""):
            return mult
    return 1.0


def strip_scale(text: str) -> str:
    """The header without its scale marker: "Revenue ($000)" → "Revenue"."""
    for rx, _ in _SCALE_RES:
        text = rx.sub(" ", text)
    return text


def _pct_like(name: str, present: list[Any]) -> bool:
    marked = sum(1 for v in present if isinstance(v, str) and v.strip().endswith("%"))
    return marked >= max(1, len(present) / 2) or bool(re.search(r"%|\bpercent|\(pct\)", name or "", re.I))


def _profile(name: str, values: list[Any], table_scale: float = 1.0) -> ColumnProfile:
    present = [v for v in values if not _blank(v)]
    numeric = _share(values, lambda v: parse_number(v) is not None)
    if numeric >= 0.8:
        dtype = "number"
    elif period_reader(name, values):
        dtype = "period"
    else:
        dtype = "text"
    samples = [str(v) for v in present[:3]]
    null_rate = 1 - len(present) / len(values) if values else 1.0
    prof = ColumnProfile(
        name, dtype, samples, null_rate, column_currency_of(name, present) if dtype == "number" else None
    )
    if dtype == "number":
        if _pct_like(name, present):
            prof.unit_hint = "pct"
        else:
            own = scale_of(name)
            prof.scale = own if own != 1.0 else (table_scale if not _NOT_MONEY.search(name) else 1.0)
    else:
        prof.unit_hint = "date" if dtype == "period" else "text"
    return prof


def column_currency_of(name: str, present: list[Any]) -> str | None:
    """The header's currency ("Revenue (USD)"), else the one most cells carry ("US$1,200")."""
    head = currency_in(name)
    if head:
        return head
    marked = [c for v in present if isinstance(v, str) and (c := currency_in(v))]
    if marked and len(marked) >= len(present) / 2 and len(set(marked)) == 1:
        return marked[0]
    return None


def _granularity(keys: list[str]) -> str:
    if not keys:
        return "none"
    if all("-Q" in k for k in keys):
        return "quarter"
    if all(re.match(r"^\d{4}-\d{2}$", k) for k in keys):
        return "month"
    if all(re.match(r"^\d{4}$", k) for k in keys):
        return "year"
    return "none"


def _dominant(grans: list[str]) -> str:
    """The granularity most periods share; the finer one on a tie."""
    c = Counter(grans)
    return max(c, key=lambda g: (c[g], {"month": 3, "quarter": 2, "year": 1}.get(g, 0)))


def _unique(names: list[str]) -> list[str]:
    seen: Counter = Counter()
    out = []
    for n in names:
        seen[n] += 1
        out.append(n if seen[n] == 1 else f"{n} ({seen[n]})")
    return out


_YEAR_HEAD = re.compile(r"^\s*(fiscal |calendar )?(year|yr|fy)\s*$", re.I)
_MONTH_HEAD = re.compile(r"^\s*(month|mon|mth|period)\s*$", re.I)
_QUARTER_HEAD = re.compile(r"^\s*(quarter|qtr|q)\s*$", re.I)


def _combine_split_dates(
    header: list[str], rows: list[list[Any]]
) -> tuple[list[str], list[list[Any]], str | None]:
    """Year | Month (or Year | Quarter) columns become one period column."""
    cols = list(zip(*rows, strict=False)) if rows else []
    if not cols:
        return header, rows, None

    def years(i: int) -> bool:
        return (
            _share(
                list(cols[i]),
                lambda v: (n := parse_number(v)) is not None and n == int(n) and 1990 <= n <= 2100,
            )
            >= 0.8
        )

    def month_of(v: Any) -> int | None:
        n = parse_number(v)
        if n is not None:
            return int(n) if n == int(n) and 1 <= n <= 12 else None
        return _month_word(str(v).strip())

    def quarter_of_cell(v: Any) -> int | None:
        m = re.fullmatch(r"q?([1-4])", str(v).strip().lower().split(".")[0])
        return int(m.group(1)) if m else None

    yi = next(
        (i for i, h in enumerate(header) if i < len(cols) and _YEAR_HEAD.match(h or "") and years(i)), None
    )
    if yi is None:
        return header, rows, None
    for rx, cell, fmt in (
        (_MONTH_HEAD, month_of, "{y:04d}-{p:02d}"),
        (_QUARTER_HEAD, quarter_of_cell, "{y:04d}-Q{p}"),
    ):
        pi = next(
            (
                i
                for i, h in enumerate(header)
                if i != yi
                and i < len(cols)
                and rx.match(h or "")
                and _share(list(cols[i]), lambda v, f=cell: f(v) is not None) >= 0.8
            ),
            None,
        )
        if pi is None:
            continue
        keep = min(yi, pi)
        new_header = [("Period" if i == keep else h) for i, h in enumerate(header) if i != max(yi, pi)]
        new_rows = []
        for r in rows:
            y, p = parse_number(r[yi]), cell(r[pi])
            key = fmt.format(y=int(y), p=p) if y is not None and p is not None else None
            new_rows.append([(key if i == keep else v) for i, v in enumerate(r) if i != max(yi, pi)])
        return new_header, new_rows, f"{header[yi]} and {header[pi]} columns combined into one date"
    return header, rows, None


_STOCK = re.compile(
    r"cash|balance|headcount|head count|employees|\bfte|staff|customers|clients|subscribers|users|"
    r"accounts|logos|\bmrr\b|\barr\b|run.?rate|inventory|pipeline|backlog|ending|closing|end of|outstanding",
    re.I,
)
_FLOW = re.compile(r"\bnew\b|added|churn|lost|gained|acquired|won|signups?|sign-ups?", re.I)
# A dimension column of ledger lines; "Account" or "Name" alone could be customers, so their values decide.
_CATEGORY_HEAD = re.compile(
    r"category|line ?item|metric|\bgl\b|ledger|\bclass\b|cost cent(?:re|er)|department|"
    r"expense type|account (?:name|type|category)|chart of accounts",
    re.I,
)
_CUSTOMER_HEAD = re.compile(r"customer|client|company|logo|tenant|subscriber|buyer|\borg", re.I)
_AMOUNT_HEAD = re.compile(r"amount|value|total|net|balance|sum|usd|cad|eur|gbp|\$", re.I)
_FORECAST = re.compile(
    r"\b(budget|bud|plan|planned|forecast|fcst|fcast|projected|projection|proj|target|estimate|est)\b", re.I
)


def is_stock(name: str) -> bool:
    """A balance (cash, customers, headcount): several rows in a period keep the last, not the sum."""
    return bool(_STOCK.search(name)) and not _FLOW.search(name)


def _line_items(values: list[str]) -> bool:
    """Most distinct values are financial line names (Revenue, Hosting, Rent), not customers."""
    from vik.model.mapping import _rule_match  # mapping imports this module

    distinct = sorted({v for v in values if v})
    if not distinct or len(distinct) > 120:
        return False
    return sum(1 for v in distinct if _rule_match(v)[0] not in (None, "period")) >= 0.5 * len(distinct)


def _forced(header: list[str], cols: list, column: str) -> tuple[int, Reader]:
    if column not in header:
        raise ValueError(f"No column named “{column}”.")
    i = header.index(column)
    found = period_reader(column, list(cols[i]) if i < len(cols) else [], threshold=0.5)
    if not found:
        sample = next((str(v) for v in (cols[i] if i < len(cols) else []) if not _blank(v)), "")
        raise ValueError(
            f"“{column}” doesn't read as dates"
            + (f" (it has “{sample[:40]}”)" if sample else "")
            + ". Pick the column that holds the dates, or keep this one as a raw column."
        )
    return i, found[0]


def parse_table(t: RawTable, period_column: str | None = None) -> ParsedDataset:
    """`period_column` makes that column the date axis (the founder chose it in the mapping)."""
    header, rows = t.header, t.rows
    notes: list[str] = []
    # Drop columns with no header and no values (trailing commas, spacer columns in spreadsheets).
    width = max([len(header), *(len(r) for r in rows)]) if rows else len(header)
    padded = [list(r) + [None] * (width - len(r)) for r in rows]
    header = list(header) + [""] * (width - len(header))
    keep = [i for i in range(width) if not _blank(header[i]) or any(not _blank(r[i]) for r in padded)]
    header = [str(header[i] or "").strip() or f"col_{i + 1}" for i in keep]
    rows = [[r[i] for i in keep] for r in padded]
    header, rows, combined = _combine_split_dates(header, rows)
    if combined:
        notes.append(combined)
    header = _unique(header)
    cols = list(zip(*rows, strict=False)) if rows else [[] for _ in header]
    table_scale = scale_of(t.preamble)

    period_idx: int | None = None
    reader: Reader = parse_period
    if period_column is not None:
        period_idx, reader = _forced(header, cols, period_column)
    else:
        found: list[tuple[int, Reader]] = []
        for i, c in enumerate(cols):
            r = period_reader(header[i], list(c))
            if r and (
                r[1]
                or _share(list(c), lambda v: parse_number(v) is not None and parse_period(v) is None) < 0.5
            ):
                found.append((i, r[0]))
        if found:
            period_idx, reader = found[0]
        if len(found) > 1:
            # Two period columns means a cohort table: the cohort/signup column is the dimension. Otherwise the
            # finest one is the axis (a Month column beside a Year column).
            cohort_words = re.compile(r"cohort|signup|sign_up|start|acquired|joined|first", re.I)
            non_cohort = [f for f in found if not cohort_words.search(header[f[0]])] or [found[-1]]
            rank = {"month": 3, "quarter": 2, "year": 1}

            def fineness(f: tuple[int, Reader]) -> int:
                got = [g[1] for v in cols[f[0]] if not _blank(v) and (g := f[1](v))]
                return rank.get(_dominant(got), 0) if got else 0

            period_idx, reader = max(non_cohort, key=lambda f: (fineness(f), -f[0]))

    if period_idx is None:
        wide = _wide(t, header, rows, notes, table_scale)
        if wide:
            return wide
        columns = [
            _profile(h, list(cols[i]) if i < len(cols) else [], table_scale) for i, h in enumerate(header)
        ]
        raw_rows = [
            {
                h: (parse_number(r[i]) if parse_number(r[i]) is not None else r[i])
                for i, h in enumerate(header)
            }
            for r in rows
        ]
        for c in columns:
            c.scale = 1.0  # raw rows keep the values as printed
        notes.append(
            "No date axis found; kept as a raw table"
            + ("" if not rows else ". Its figures are read as facts")
        )
        return ParsedDataset(
            t.name, "raw", None, "none", False, columns, raw_rows, fingerprint(header), notes
        )

    profiles = [_profile(h, list(cols[i]), table_scale) for i, h in enumerate(header)]
    profiles[period_idx].dtype, profiles[period_idx].unit_hint, profiles[period_idx].scale = (
        "period",
        "date",
        1.0,
    )
    scales = [p.scale for p in profiles]
    _note_scale(profiles, notes)

    def num(i: int, v: Any) -> float | None:
        n = parse_number(v)
        return n * scales[i] if n is not None and scales[i] != 1.0 else n

    # Rows whose period is coarser or finer than the table's (a "2024" subtotal in a monthly table) are left out.
    # Each row keeps its line in the file (`t.lines`; a table built in code counts from the header on line 1).
    lines = t.lines if len(t.lines) == len(rows) else list(range(2, len(rows) + 2))
    keyed = [(reader(r[period_idx]), r, n) for r, n in zip(rows, lines, strict=True)]
    keyed = [(p, r, n) for p, r, n in keyed if p]
    if not keyed:
        return ParsedDataset(
            t.name,
            "raw",
            None,
            "none",
            False,
            profiles,
            [],
            fingerprint(header),
            [*notes, "No rows with a readable date"],
        )
    gran = _dominant([p[1] for p, _, _ in keyed])
    off = [p[0] for p, _, _ in keyed if p[1] != gran]
    if off:
        keyed = [(p, r, n) for p, r, n in keyed if p[1] == gran]
        notes.append(
            f"Left out {len(off)} subtotal row{'s' if len(off) > 1 else ''} ({', '.join(sorted(set(off))[:3])})"
        )
    # Oldest first within each period, whatever the file order (bank exports list the newest first), so a
    # balance takes the value on the period's latest date.
    first_day = _dayfirst([v for v in cols[period_idx] if not _blank(v)])
    days = [_day_of(r[period_idx], first_day) for _, r, _ in keyed]
    if any(days):
        order = sorted(
            range(len(keyed)), key=lambda j: (period_sort_key(keyed[j][0][0]), days[j] or date.min)
        )
        keyed = [keyed[j] for j in order]
    text_dims = [i for i, p in enumerate(profiles) if i != period_idx and p.dtype in ("text", "period")]
    numeric_cols = [i for i, p in enumerate(profiles) if p.dtype == "number"]

    # A long customer/cohort table repeats each period once per customer; a P&L with a notes
    # column has one row per period and must stay a series.
    repeats = len(keyed) > len({p[0] for p, _, _ in keyed})
    if text_dims and numeric_cols and repeats:
        dim_idx = text_dims[0]
        dim_values = [str(r[dim_idx]).strip() for _, r, _ in keyed if not _blank(r[dim_idx])]
        is_cohort = _share(list(cols[dim_idx]), lambda v: parse_period(v) is not None) >= 0.8
        ledger = (
            not is_cohort
            and not _CUSTOMER_HEAD.search(header[dim_idx])
            and (_CATEGORY_HEAD.search(header[dim_idx]) or _line_items(dim_values))
        )
        if ledger:
            return _pivot(t, header, keyed, dim_idx, numeric_cols, num, scales, notes)
        kind = "cohort_mrr" if is_cohort else "customer_mrr"
        out_rows = []
        for p, r, _ in keyed:
            dim_raw = r[dim_idx]
            dim = parse_period(dim_raw)[0] if is_cohort and parse_period(dim_raw) else str(dim_raw).strip()
            row = {"period": p[0], "dim": dim}
            for i in numeric_cols:
                row[header[i]] = num(i, r[i])
            out_rows.append(row)
        keys = sorted({r["period"] for r in out_rows}, key=period_sort_key)
        return ParsedDataset(
            t.name,
            kind,
            header[period_idx],
            _granularity(keys),
            False,
            profiles,
            out_rows,
            fingerprint(header),
            notes,
        )

    merged: dict[str, dict[str, Any]] = {}
    stock = [is_stock(h) for h in header]
    for p, r, n in keyed:
        row = merged.get(p[0])
        first = row is None
        if first:
            row = merged[p[0]] = {"period": p[0], SRC_ROWS: []}
        row[SRC_ROWS].append(n)
        for i, h in enumerate(header):
            if i == period_idx:
                continue
            v = num(i, r[i])
            if first or row.get(h) is None or stock[i]:
                if v is not None or first:
                    row[h] = v
            elif v is not None:
                row[h] += v
    if repeats:
        notes.append(
            f"Several rows per {gran} (e.g. daily or weekly data): amounts are summed per {gran}; "
            "balances (cash, customers, headcount) take the period's last value"
        )
    out_rows = sorted(merged.values(), key=lambda r: period_sort_key(r["period"]))
    keys = [r["period"] for r in out_rows]
    return ParsedDataset(
        t.name,
        "series",
        header[period_idx],
        _granularity(keys),
        False,
        profiles,
        out_rows,
        fingerprint(header),
        notes,
    )


def _note_scale(profiles: list[ColumnProfile], notes: list[str]) -> None:
    scaled = [p for p in profiles if p.scale != 1.0]
    if scaled:
        word = {1e3: "thousands", 1e6: "millions", 1e9: "billions"}
        by: dict[float, list[str]] = {}
        for p in scaled:
            by.setdefault(p.scale, []).append(p.name)
        for s, names in by.items():
            shown = ", ".join(names[:3]) + (f" and {len(names) - 3} more" if len(names) > 3 else "")
            notes.append(f"Values in {word.get(s, f'×{s:,.0f}')} ({shown}): multiplied by {s:,.0f}")


def _pivot(
    t: RawTable,
    header: list[str],
    keyed: list,
    dim_idx: int,
    numeric_cols: list[int],
    num,
    scales: list[float],
    notes: list[str],
) -> ParsedDataset:
    """A ledger (date | account | amount, one line per account per period) becomes one column per account."""
    val_idx = next((i for i in numeric_cols if _AMOUNT_HEAD.search(header[i])), numeric_cols[0])
    merged: dict[str, dict[str, Any]] = {}
    raw_by: dict[str, list[Any]] = {}
    for p, r, _ in keyed:
        label = str(r[dim_idx]).strip() if not _blank(r[dim_idx]) else ""
        v = num(val_idx, r[val_idx])
        if not label:
            continue
        row = merged.setdefault(p[0], {"period": p[0]})
        raw_by.setdefault(label, []).append(r[val_idx])
        if v is None:
            row.setdefault(label, None)
        elif row.get(label) is None or is_stock(label):
            row[label] = v
        else:
            row[label] += v
    labels = list(raw_by)
    out_rows = sorted(merged.values(), key=lambda r: period_sort_key(r["period"]))
    for row in out_rows:
        for lb in labels:
            row.setdefault(lb, None)
    keys = [r["period"] for r in out_rows]
    notes.append(
        f"One row per {header[dim_idx]} per period: each {header[dim_idx]} became a column of "
        f"{header[val_idx]}, summed per period"
    )
    columns = [ColumnProfile("period", "period", keys[:3], 0.0, unit_hint="date")]
    for lb in labels:
        prof = _profile(lb, raw_by[lb])
        prof.scale = 1.0 if prof.unit_hint == "pct" else scales[val_idx]
        columns.append(prof)
    return ParsedDataset(
        t.name, "series", "period", _granularity(keys), True, columns, out_rows, fingerprint(labels), notes
    )


def _wide(
    t: RawTable, header: list[str], rows: list[list[Any]], notes: list[str], table_scale: float
) -> ParsedDataset | None:
    """Periods across the header row: transpose so each line item is a column."""
    hp = [None if _FORECAST.search(h) else parse_period(h) for h in header]
    first = next((j for j, p in enumerate(hp) if p), None)
    if first is None or first == 0:
        return None
    # The label column: the most text-like column left of the first period (an account-code column is skipped).
    left = range(first)
    label_idx = max(left, key=lambda i: (_share([r[i] for r in rows], lambda v: parse_number(v) is None), -i))
    right = list(range(first, len(header)))
    pcols = [j for j in right if hp[j]]
    if len(pcols) < 2 or len(pcols) < 0.5 * len(right):
        return None
    gran = _dominant([hp[j][1] for j in pcols])
    dropped = [header[j] for j in pcols if hp[j][1] != gran]
    pcols = [j for j in pcols if hp[j][1] == gran]
    seen: set[str] = set()
    dupes = [j for j in pcols if hp[j][0] in seen or seen.add(hp[j][0])]
    pcols = [j for j in pcols if j not in dupes]
    if len(pcols) < 2:
        return None
    notes.append("Transposed: periods were columns")
    skipped = dropped + [header[j] for j in right if not hp[j] and not _blank(header[j])]
    if skipped:
        notes.append(
            f"Left out {len(skipped)} column{'s' if len(skipped) > 1 else ''} that aren't single "
            f"{gran}s ({', '.join(skipped[:3])}{'…' if len(skipped) > 3 else ''})"
        )
    if dupes:
        notes.append(f"Left out {len(dupes)} repeated period column{'s' if len(dupes) > 1 else ''}")

    labels: list[str] = []
    values: list[list[Any]] = []
    section = ""
    used: set[str] = set()
    for r in rows:
        label = str(r[label_idx]).strip() if not _blank(r[label_idx]) else ""
        vals = [r[j] for j in pcols]
        if not label:
            continue
        if all(parse_number(v) is None for v in vals) and all(_blank(v) for v in vals):
            section = label  # a heading row ("Operating expenses") with no figures
            continue
        name = label
        if name in used:
            name = f"{section} · {label}" if section and f"{section} · {label}" not in used else label
            k = 2
            while name in used:
                name, k = f"{label} ({k})", k + 1
        used.add(name)
        labels.append(name)
        values.append(vals)
    periods = [hp[j][0] for j in pcols]
    columns = [_profile(lb, vals, table_scale) for lb, vals in zip(labels, values, strict=True)]
    order = sorted(range(len(periods)), key=lambda j: period_sort_key(periods[j]))
    out_rows = []
    for j in order:
        row: dict[str, Any] = {"period": periods[j]}
        for c, vals in zip(columns, values, strict=True):
            n = parse_number(vals[j])
            row[c.name] = n * c.scale if n is not None else None
        out_rows.append(row)
    _note_scale(columns, notes)
    keys = [r["period"] for r in out_rows]
    return ParsedDataset(
        t.name,
        "series",
        header[label_idx] or "transposed",
        _granularity(keys),
        True,
        [ColumnProfile("period", "period", keys[:3], 0.0, unit_hint="date"), *columns],
        out_rows,
        fingerprint(labels),
        notes,
    )


def parse_upload(
    data: bytes, filename: str, period_column: dict[str, str] | None = None
) -> list[ParsedDataset]:
    """`period_column` maps a table name to the column the founder chose as its dates."""
    pc = period_column or {}
    return [parse_table(t, pc.get(t.name)) for t in read_tables(data, filename) if t.header]
