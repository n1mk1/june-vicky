"""Parse CSV/XLSX in any layout into column-profiled datasets.

Handles currency/sign conventions (`$1,234`, `(5,000)`, `-$5`, `12%`), detects the period axis
(a period column, or periods across the header row, which is transposed), and recognises
dimension tables keyed by period + customer or period + signup cohort.
"""

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from typing import Any

MONTHS = {m: i + 1 for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
)}
MONTHS.update({"sept": 9, "june": 6, "july": 7})

_NUM_RE = re.compile(r"^[+-]?\d+(\.\d+)?([eE][+-]?\d+)?$")
_MISSING = {"", "-", "—", "–", "n/a", "na", "nan", "null", "none", "#n/a"}


def parse_number(raw: Any) -> float | None:
    """Normalise a spreadsheet cell to a float; None when it isn't a number."""
    if raw is None:
        return None
    if isinstance(raw, bool):
        return None
    if isinstance(raw, int | float):
        return float(raw)
    s = str(raw).strip()
    if s.lower() in _MISSING:
        return None
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg, s = True, s[1:-1].strip()
    s = s.replace("−", "-")
    if s.startswith("-"):
        neg, s = (not neg), s[1:].strip()
    elif s.startswith("+"):
        s = s[1:].strip()
    s = re.sub(r"^(cad|usd|gbp|eur|aud|c\$|us\$|a\$)\s*", "", s, flags=re.I)
    s = s.lstrip("$£€").strip()
    if s.startswith("-"):
        neg, s = (not neg), s[1:].strip()
    pct = s.endswith("%")
    if pct:
        s = s[:-1].strip()
    mult = 1.0
    m = re.match(r"^(.*?)([kmb])$", s, flags=re.I)
    if m and _NUM_RE.match(m.group(1).replace(",", "")):
        s, mult = m.group(1), {"k": 1e3, "m": 1e6, "b": 1e9}[m.group(2).lower()]
    s = s.replace(",", "").replace(" ", "").replace(" ", "")
    if not _NUM_RE.match(s):
        return None
    v = float(s) * mult
    if pct:
        v /= 100.0
    return -v if neg else v


def parse_period(raw: Any) -> tuple[str, str] | None:
    """Return (normalised key, granularity) for a period label, else None.

    Months → "YYYY-MM"; quarters → "YYYY-Qn"; years → "YYYY".
    """
    if raw is None:
        return None
    if hasattr(raw, "year") and hasattr(raw, "month"):
        return f"{raw.year:04d}-{raw.month:02d}", "month"
    s = str(raw).strip().lower().replace("_", " ")
    if not s:
        return None
    m = re.match(r"^(\d{4})[-/.](\d{1,2})(?:[-/.]\d{1,2})?(?:[ t].*)?$", s)
    if m and 1 <= int(m.group(2)) <= 12:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}", "month"
    m = re.match(r"^(\d{1,2})[-/.](\d{4})$", s)
    if m and 1 <= int(m.group(1)) <= 12:
        return f"{int(m.group(2)):04d}-{int(m.group(1)):02d}", "month"
    m = re.match(r"^([a-z]{3,9})\.?[\s\-']*(\d{2}|\d{4})$", s)
    if m and m.group(1)[:3] in MONTHS and (m.group(1) in MONTHS or m.group(1)[:3] in MONTHS):
        year = int(m.group(2))
        year = year + 2000 if year < 100 else year
        return f"{year:04d}-{MONTHS[m.group(1)[:3]]:02d}", "month"
    m = re.match(r"^(?:fy)?\s*q([1-4])[\s\-']*(?:fy)?\s*(\d{2}|\d{4})$", s)
    if m:
        year = int(m.group(2))
        year = year + 2000 if year < 100 else year
        return f"{year:04d}-Q{m.group(1)}", "quarter"
    m = re.match(r"^(?:fy)?\s*(\d{2}|\d{4})[\s\-']*q([1-4])$", s)
    if m:
        year = int(m.group(1))
        year = year + 2000 if year < 100 else year
        return f"{year:04d}-Q{m.group(2)}", "quarter"
    m = re.match(r"^(?:fy\s*)?(\d{4})$", s)
    if m and 1990 <= int(m.group(1)) <= 2100:
        return m.group(1), "year"
    return None


def period_sort_key(p: str) -> tuple[int, int]:
    if "-Q" in p:
        y, q = p.split("-Q")
        return int(y), int(q) * 3
    if "-" in p:
        y, mo = p.split("-")
        return int(y), int(mo)
    return int(p), 12


def quarter_of(month_key: str) -> str:
    y, m = month_key.split("-")
    return f"{y}-Q{(int(m) - 1) // 3 + 1}"


@dataclass
class RawTable:
    name: str
    header: list[str]
    rows: list[list[Any]]


@dataclass
class ColumnProfile:
    name: str
    dtype: str
    sample_values: list[str]
    null_rate: float

    def as_dict(self) -> dict:
        return {"name": self.name, "dtype": self.dtype, "sample_values": self.sample_values,
                "null_rate": round(self.null_rate, 3)}


@dataclass
class ParsedDataset:
    name: str
    table_kind: str
    period_axis: str | None
    granularity: str
    transposed: bool
    columns: list[ColumnProfile]
    rows: list[dict]
    fingerprint: str
    notes: list[str] = field(default_factory=list)

    def column_names(self) -> list[str]:
        return [c.name for c in self.columns]


def read_tables(data: bytes, filename: str) -> list[RawTable]:
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(data, filename)
    return [_read_csv(data, filename)]


def _read_csv(data: bytes, filename: str) -> RawTable:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    if not rows:
        return RawTable(filename, [], [])
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    return RawTable(_stem(filename), [h.strip() for h in rows[0]], rows[1:])


def _read_xlsx(data: bytes, filename: str) -> list[RawTable]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    tables: list[RawTable] = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c not in (None, "") for c in r)]
        if len(rows) < 2:
            continue
        width = max(len(r) for r in rows)
        rows = [r + [None] * (width - len(r)) for r in rows]
        header = [str(h).strip() if h is not None else f"col_{i + 1}" for i, h in enumerate(rows[0])]
        tables.append(RawTable(f"{_stem(filename)} · {ws.title}", header, rows[1:]))
    return tables


def _stem(filename: str) -> str:
    base = re.split(r"[\\/]", filename)[-1]
    return base.rsplit(".", 1)[0] if "." in base else base


def fingerprint(columns: list[str]) -> str:
    norm = "|".join(sorted(re.sub(r"[^a-z0-9]", "", c.lower()) for c in columns))
    return hashlib.sha256(norm.encode()).hexdigest()[:16]


def _share(values: list[Any], pred) -> float:
    vals = [v for v in values if v not in (None, "") and str(v).strip() != ""]
    if not vals:
        return 0.0
    return sum(1 for v in vals if pred(v)) / len(vals)


def _profile(name: str, values: list[Any]) -> ColumnProfile:
    present = [v for v in values if v not in (None, "") and str(v).strip() != ""]
    numeric = _share(values, lambda v: parse_number(v) is not None)
    dtype = "number" if numeric >= 0.8 else "period" if _share(values, lambda v: parse_period(v) is not None) >= 0.8 else "text"
    samples = [str(v) for v in present[:3]]
    null_rate = 1 - len(present) / len(values) if values else 1.0
    return ColumnProfile(name, dtype, samples, null_rate)


def _granularity(keys: list[str]) -> str:
    if not keys:
        return "none"
    if all("-Q" in k for k in keys):
        return "quarter"
    if all(re.match(r"^\d{4}-\d{2}$", k) for k in keys):
        return "month"
    if all(re.match(r"^\d{4}$", k) for k in keys):
        return "year"
    return "none"


def parse_table(t: RawTable) -> ParsedDataset:
    header, rows = t.header, t.rows
    notes: list[str] = []
    # Drop columns with no header and no values (trailing commas, spacer columns in spreadsheets).
    width = max([len(header), *(len(r) for r in rows)]) if rows else len(header)
    padded = [list(r) + [None] * (width - len(r)) for r in rows]
    header = list(header) + [""] * (width - len(header))
    keep = [i for i in range(width)
            if str(header[i] or "").strip() or any(str(r[i] if r[i] is not None else "").strip() for r in padded)]
    if len(keep) < width:
        header = [header[i] for i in keep]
        rows = [[r[i] for i in keep] for r in padded]
    cols = list(zip(*rows, strict=False)) if rows else [[] for _ in header]

    period_like = [
        i for i, c in enumerate(cols)
        if _share(list(c), lambda v: parse_period(v) is not None) >= 0.8
        and _share(list(c), lambda v: parse_number(v) is not None and parse_period(v) is None) < 0.5
    ]
    period_idx = period_like[0] if period_like else None
    if len(period_like) > 1:
        # Two period columns means a cohort table: the cohort/signup column is the dimension.
        cohort_words = re.compile(r"cohort|signup|sign_up|start|acquired|joined|first", re.I)
        non_cohort = [i for i in period_like if not cohort_words.search(header[i])]
        period_idx = non_cohort[0] if non_cohort else period_like[-1]
    header_periods = [parse_period(h) for h in header[1:]]
    wide = period_idx is None and header_periods and sum(p is not None for p in header_periods) >= max(2, 0.6 * len(header_periods))

    if wide:
        notes.append("Transposed: periods were columns")
        labels = [str(r[0]).strip() for r in rows]
        periods = [p[0] if p else None for p in header_periods]
        out_rows = []
        for j, p in enumerate(periods):
            if p is None:
                continue
            row: dict[str, Any] = {"period": p}
            for i, label in enumerate(labels):
                if label:
                    row[label] = parse_number(rows[i][j + 1])
            out_rows.append(row)
        out_rows.sort(key=lambda r: period_sort_key(r["period"]))
        columns = [_profile(label, [rows[i][j + 1] for j in range(len(periods))]) for i, label in enumerate(labels) if label]
        keys = [r["period"] for r in out_rows]
        return ParsedDataset(t.name, "series", header[0] or "transposed", _granularity(keys), True,
                             [ColumnProfile("period", "period", keys[:3], 0.0), *columns], out_rows,
                             fingerprint(labels), notes)

    if period_idx is None:
        columns = [_profile(h, list(cols[i]) if i < len(cols) else []) for i, h in enumerate(header)]
        raw_rows = [{h: (parse_number(r[i]) if parse_number(r[i]) is not None else r[i]) for i, h in enumerate(header)} for r in rows]
        notes.append("No period axis found; kept as a raw table")
        return ParsedDataset(t.name, "raw", None, "none", False, columns, raw_rows, fingerprint(header), notes)

    profiles = [_profile(h, list(cols[i])) for i, h in enumerate(header)]
    profiles[period_idx].dtype = "period"
    text_dims = [i for i, p in enumerate(profiles) if i != period_idx and p.dtype in ("text", "period")]
    numeric_cols = [i for i, p in enumerate(profiles) if p.dtype == "number"]

    # A long customer/cohort table repeats each period once per customer; a P&L with a notes
    # column has one row per period and must stay a series.
    period_keys = [parse_period(r[period_idx]) for r in rows]
    repeats = len([p for p in period_keys if p]) > len({p[0] for p in period_keys if p})
    if text_dims and numeric_cols and repeats:
        dim_idx = text_dims[0]
        is_cohort = _share(list(cols[dim_idx]), lambda v: parse_period(v) is not None) >= 0.8
        kind = "cohort_mrr" if is_cohort else "customer_mrr"
        out_rows = []
        for r in rows:
            p = parse_period(r[period_idx])
            if not p:
                continue
            dim_raw = r[dim_idx]
            dim = parse_period(dim_raw)[0] if is_cohort and parse_period(dim_raw) else str(dim_raw).strip()
            row = {"period": p[0], "dim": dim}
            for i in numeric_cols:
                row[header[i]] = parse_number(r[i])
            out_rows.append(row)
        keys = sorted({r["period"] for r in out_rows}, key=period_sort_key)
        return ParsedDataset(t.name, kind, header[period_idx], _granularity(keys), False, profiles, out_rows,
                             fingerprint(header), notes)

    merged: dict[str, dict[str, Any]] = {}
    for r in rows:
        p = parse_period(r[period_idx])
        if not p:
            continue
        row = merged.setdefault(p[0], {"period": p[0]})
        for i, h in enumerate(header):
            if i == period_idx:
                continue
            row[h] = parse_number(r[i])
    out_rows = sorted(merged.values(), key=lambda r: period_sort_key(r["period"]))
    keys = [r["period"] for r in out_rows]
    return ParsedDataset(t.name, "series", header[period_idx], _granularity(keys), False, profiles, out_rows,
                         fingerprint(header), notes)


def parse_upload(data: bytes, filename: str) -> list[ParsedDataset]:
    return [parse_table(t) for t in read_tables(data, filename) if t.header]
