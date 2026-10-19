"""Regression tests for the metrics-engine findings of the 2026-10-18 QA sweep (its report is in git history,
section 3.1). Each test reproduces the report's repro; the test name carries the finding id."""

import pytest

from vik.model.context import build_context, sample_context
from vik.model.figures import Fact, Figure, MissingFigure, Source
from vik.model.financials import DatasetInput
from vik.model.fx import Rates
from vik.model.loader import parse_number, parse_upload, scale_of
from vik.model.mapping import _rule_match, propose_heuristic
from vik.model.metrics import facts_by_metric
from vik.model.truth import check_claim, quick_claims


def inputs_of(*files: tuple[str, str], currency: str = "cad") -> list[DatasetInput]:
    """Parse CSV texts and map them with the heuristic mapper, as an upload would be before confirmation."""
    out = []
    for text, name in files:
        for ds in parse_upload(text.encode(), name):
            out.append(
                DatasetInput(
                    id=f"{name}#{ds.name}",
                    name=ds.name,
                    table_kind=ds.table_kind,
                    granularity=ds.granularity,
                    transposed=ds.transposed,
                    rows=ds.rows,
                    mapping=[m.model_dump() for m in propose_heuristic(ds, currency)],
                )
            )
    return out


def ctx_of(
    *files: tuple[str, str],
    currency: str = "cad",
    facts: list[Fact] | None = None,
    fx: Rates | None = None,
    **profile,
):
    return build_context(
        inputs_of(*files, currency=currency), facts or [], {"currency": currency, **profile}, fx=fx
    )


def fact(
    name: str,
    value: float,
    unit: str = "cad",
    kind: str = "historical",
    period: str | None = None,
    status: str = "confirmed",
    **kw,
) -> Fact:
    return Fact(
        id=f"f-{name}-{period}-{status}",
        name=name,
        value=value,
        unit=unit,
        kind=kind,
        period=period,
        status=status,
        source=Source(kind="pdf", ref="deck.pdf p.2"),
        **kw,
    )


def months(start_year: int, start_month: int, n: int) -> list[str]:
    out = []
    y, m = start_year, start_month
    for _ in range(n):
        out.append(f"{y:04d}-{m:02d}")
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


# ---------------------------------------------------------------- MET-2

PROFITABLE = "Month,Revenue,Net income\n" + "".join(f"{p},10000,2000\n" for p in months(2025, 1, 6))


@pytest.mark.parametrize("plan", [5000.0, 0.0])
def test_met2_forecast_fact_against_negative_history_does_not_crash(plan):
    ctx = ctx_of((PROFITABLE, "pl.csv"), facts=[fact("burn", plan, kind="forecast")])
    assert ctx.figure("burn").value == pytest.approx(-2000)
    gap = next(a for a in ctx.attack if a.id == "gap_burn")
    assert gap.severity >= 0.9  # a sign flip against history is the largest gap


# ---------------------------------------------------------------- MET-3

NEG_COSTS = (
    "Month,Revenue,COGS,Sales and marketing,R&D,G&A,Cash\n"
    "2025-01,10000,(3000),(4000),(8000),(2000),500000\n"
    "2025-02,10000,(3000),(4000),(8000),(2000),493000\n"
)


def test_met3_costs_shown_as_negatives_are_read_as_costs():
    ctx = ctx_of((NEG_COSTS, "pl.csv"))
    assert ctx.figure("gross_margin").value == pytest.approx(0.7)
    assert ctx.model.value("net_income", "2025-02") == pytest.approx(-7000)
    assert ctx.figure("burn").value == pytest.approx(7000)
    assert ctx.figure("runway_months").value == pytest.approx(493000 / 7000)
    assert ctx.figure("opex_share_r_and_d").value == pytest.approx(8 / 14)
    assert any("negatives" in n for n in ctx.model.notes)


def test_met3_a_single_credit_in_a_cost_column_is_not_flipped():
    csv = "Month,Revenue,COGS\n2025-01,10000,3000\n2025-02,10000,3000\n2025-03,10000,-500\n"
    ctx = ctx_of((csv, "pl.csv"))
    assert ctx.model.value("cogs", "2025-01") == 3000 and ctx.model.value("cogs", "2025-03") == -500
    assert not ctx.model.notes


# ---------------------------------------------------------------- MET-4


def test_met4_day_first_first_of_month_dates_keep_twelve_months():
    csv = "Date,Revenue\n" + "".join(f"01/{m:02d}/2025,{m * 1000}\n" for m in range(1, 13))
    [ds] = parse_upload(csv.encode(), "t.csv")
    assert [r["period"] for r in ds.rows] == months(2025, 1, 12)
    assert ds.rows[-1]["Revenue"] == 12000


def test_met4_month_first_first_of_month_dates_still_read_month_first():
    csv = "Date,Revenue\n" + "".join(f"{m:02d}/01/2025,{m * 1000}\n" for m in range(1, 13))
    [ds] = parse_upload(csv.encode(), "t.csv")
    assert [r["period"] for r in ds.rows] == months(2025, 1, 12)


# ---------------------------------------------------------------- MET-7


def test_met7_harbor_labs_services_share_uses_the_same_quarter():
    ctx = sample_context("harbor-labs")
    assert "services_mix" not in {
        a.id for a in ctx.attack
    }  # 62,000 of 695,804 is 9%, under the 12% threshold
    f = ctx.figure("one_time_revenue_share")
    assert f.value == pytest.approx(1 - 633_804 / 695_804) and f.period == "2026-Q3"


def test_met7_services_share_is_a_citable_figure_for_the_latest_period():
    csv = "Quarter,Subscription revenue,Professional services\nQ1 2025,30000,30000\nQ2 2025,30000,30000\n"
    ctx = ctx_of((csv, "q.csv"))
    item = next(a for a in ctx.attack if a.id == "services_mix")
    assert "50%" in item.detail and "latest period" in item.detail and "month" not in item.detail
    assert "one_time_revenue_share" in item.figure_ids
    assert ctx.figure("one_time_revenue_share").display == "50%"


# ---------------------------------------------------------------- MET-8


def test_met8_a_forecast_word_elsewhere_does_not_make_a_claim_a_forecast():
    [arr] = quick_claims("Our ARR is $3M today. Next year we will hire five engineers.")
    assert arr["metric"] == "arr" and arr["is_forecast"] is False
    [plan] = quick_claims("We will reach ARR of $3M next year.")
    assert plan["is_forecast"] is True


# ---------------------------------------------------------------- MET-9


def test_met9_foreign_fact_without_a_rate_is_not_relabelled():
    ctx = build_context(
        [], [fact("mrr", 100_000, unit="usd", period="2025-01")], {"currency": "cad"}, fx=None
    )
    for fid in ("mrr", "arr"):
        f = ctx.figures[fid]
        assert isinstance(f, MissingFigure), f"{fid} shown as {f.display}"
    assert "USD→CAD" in ctx.figures["mrr"].reason
    cash = build_context(
        [], [fact("cash", 50_000, unit="usd"), fact("cac", 900, unit="eur")], {"currency": "cad"}
    )
    assert isinstance(cash.figures["cash_balance"], MissingFigure) and isinstance(
        cash.figures.get("cac"), MissingFigure | None
    )


def test_met9_foreign_fact_with_a_rate_converts():
    fx = Rates(daily={"usd": {"2025-01-15": 1.4}})
    ctx = build_context([], [fact("mrr", 100_000, unit="usd", period="2025-01")], {"currency": "cad"}, fx=fx)
    assert ctx.figure("mrr").value == pytest.approx(140_000) and ctx.figure("mrr").unit == "cad"


def test_met9_foreign_cash_burn_fact_is_not_used_in_the_burn_multiple():
    rows = months(2024, 1, 13)
    csv = "Month,Subscription revenue,Net income\n" + "".join(
        f"{p},{10000 + i * 1000},-20000\n" for i, p in enumerate(rows)
    )
    ctx = ctx_of((csv, "pl.csv"), facts=[fact("net_burn_l12m", 999_999, unit="usd")])
    bm = ctx.figure("burn_multiple")
    assert bm is not None and "999,999" not in bm.formula and "P&L basis" in bm.formula


# ---------------------------------------------------------------- MET-10


def test_met10_plan_facts_are_checked_against_history():
    ctx = sample_context("tidewell")
    c = check_claim("mrr", 129_000, "cad", ctx.figures, facts=ctx.facts)
    assert c.status == "forecast_as_fact"
    # The nearest plan (Q4 2026: $129K MRR against $107K today) is within tolerance; the 2027-28 plan points that
    # are far above today's history (a growth plan is meant to be) don't raise a gap of their own.
    ids = [a.id for a in ctx.attack]
    assert "gap_mrr" not in ids and "gap_arr" not in ids
    assert not [a for a in ctx.attack if a.id.startswith("gap_") and a.severity >= 0.95]
    assert len(ids) == len(set(ids))


SIX_MONTHS = "Month,Subscription revenue\n" + "".join(f"{p},100000\n" for p in months(2025, 1, 6))


def _plan(value: float, period: str | None) -> Fact:
    return fact("plan.mrr", value, kind="forecast", period=period)


def test_met10_the_nearest_plan_period_is_the_one_compared():
    ctx = ctx_of(
        (SIX_MONTHS, "pl.csv"),
        facts=[
            _plan(1_000_000, "2027-Q4"),
            _plan(200_000, "2025-Q3"),
            _plan(150_000, "2025-Q4"),
            _plan(50_000, "2025-Q1"),
        ],
    )
    [gap] = [a for a in ctx.attack if a.id.startswith("gap_")]
    assert (
        gap.id == "gap_mrr" and "Q3 2025" in gap.detail and "$200.0K" in gap.detail
    )  # not the stale Q1, not 2027
    assert gap.severity < 0.95


def test_met10_a_far_future_plan_alone_is_no_gap_when_the_next_step_is_on_track():
    ctx = ctx_of((SIX_MONTHS, "pl.csv"), facts=[_plan(110_000, "2025-Q3"), _plan(1_000_000, "2027-Q4")])
    assert not [a for a in ctx.attack if a.id.startswith("gap_")]
    far = ctx_of(
        (SIX_MONTHS, "pl.csv"), facts=[_plan(1_000_000, "2027-Q4")]
    )  # only a long-range target (30 months out)
    assert not [a for a in far.attack if a.id.startswith("gap_")]
    stale = ctx_of(
        (SIX_MONTHS, "pl.csv"), facts=[_plan(10_000, "2025-03")]
    )  # a plan for a month already in history
    assert not [a for a in stale.attack if a.id.startswith("gap_")]


# ---------------------------------------------------------------- MET-11


def _ten_pct(skip: str | None = None) -> str:
    rows = months(2024, 1, 14)  # Jan 24 .. Feb 25
    return "Month,Subscription revenue\n" + "".join(
        f"{p},{10000 * 1.1**i:.4f}\n" for i, p in enumerate(rows) if p != skip
    )


def test_met11_growth_windows_are_calendar_months():
    ctx = ctx_of((_ten_pct(skip="2024-12"), "pl.csv"))
    g12 = ctx.figure("cmgr_12m")
    assert g12.value == pytest.approx(0.10) and g12.period == "2024-02..2025-02"
    assert isinstance(ctx.figures["cmgr_recent"], MissingFigure)  # Dec 24 has no data
    yoy = ctx.figure("revenue_growth_yoy")
    assert yoy.value == pytest.approx(1.1**12 - 1) and "Feb 24" in yoy.formula
    assert isinstance(ctx.figures["revenue_growth"], Figure)  # Feb vs Jan 25 are both there
    nna = ctx.figure("net_new_arr_l12m")
    assert nna.value == pytest.approx((10000 * 1.1**13 - 10000 * 1.1) * 12)


def test_met11_complete_data_is_unchanged():
    ctx = ctx_of((_ten_pct(), "pl.csv"))
    assert ctx.figure("cmgr_12m").value == pytest.approx(0.10)
    assert ctx.figure("cmgr_recent").value == pytest.approx(0.10)


# ---------------------------------------------------------------- MET-12


def _credits_csv(net_income: int = -50000, credit_months: tuple[str, ...] = ("2024-06", "2024-12")) -> str:
    rows = months(2024, 1, 13)  # Jan 24 .. Jan 25
    lines = ["Month,Subscription revenue,Net income,SR&ED credits"]
    for i, p in enumerate(rows):
        credit = 60000 if p in credit_months else ""
        ni = net_income + (60000 if credit else 0)
        lines.append(f"{p},{10000 + i * 10000 / 12:.4f},{ni},{credit}")
    return "\n".join(lines) + "\n"


def test_met12_credits_in_some_months_are_excluded():
    ctx = ctx_of((_credits_csv(), "pl.csv"))
    # Window Feb 24 .. Jan 25: net loss 600,000 - 120,000 credits; net new ARR 120,000.
    assert ctx.figure("burn_multiple").value == pytest.approx(480_000 / 120_000)
    assert ctx.figure("burn_multiple_ex_credits").value == pytest.approx(600_000 / 120_000)


# ---------------------------------------------------------------- MET-13


def test_met13_exchange_rate_moves_are_not_retention():
    rows = months(2024, 1, 13)
    usd = {f"{p}-15": 1.30 + 0.15 * i / 12 for i, p in enumerate(rows)}
    csv = "Month,Customer,MRR (USD)\n" + "".join(f"{p},Acme,1000\n{p},Beta,500\n" for p in rows)
    pl = "Month,Subscription revenue\n" + "".join(f"{p},1500\n" for p in rows)
    ctx = ctx_of((csv, "c.csv"), (pl, "pl.csv"), fx=Rates(daily={"usd": usd}))
    converted = ctx.model.customer_mrr
    assert converted["2025-01"]["Acme"] == pytest.approx(1450)  # absolute money still converts
    assert (
        sum(converted["2025-01"].values()) / sum(converted["2024-01"].values()) > 1.1
    )  # the trap: rates moved 11.5%
    assert ctx.figure("nrr").value == pytest.approx(1.0)
    assert ctx.figure("grr").value == pytest.approx(1.0)


# ---------------------------------------------------------------- MET-14


def test_met14_profitable_company_has_no_burn_multiple():
    rows = months(2024, 1, 13)
    csv = "Month,Subscription revenue,Net income\n" + "".join(
        f"{p},{10000 + i * 1000},3000\n" for i, p in enumerate(rows)
    )
    ctx = ctx_of((csv, "pl.csv"))
    bm = ctx.figures["burn_multiple"]
    assert isinstance(bm, MissingFigure) and "Not burning" in bm.reason


# ---------------------------------------------------------------- MET-15 / VIZ-2


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("$(1,234.00)", -1234.0),
        (" $ (1,234.00)", -1234.0),
        ("USD (1,234)", -1234.0),
        (" $ -   ", 0.0),
        ("$ (1,234.00) ", -1234.0),
        ("(1,234.00) $", -1234.0),
        ("-", None),
        ("(5,000)", -5000.0),
        ("-$57,492.00", -57492.0),
        ("500-", -500.0),
        ("1e400", None),
        ("-1e400", None),
        (float("inf"), None),
        (float("nan"), None),
        ("nan", None),
    ],
)
def test_met15_accounting_formats_and_non_finite(raw, expected):
    assert parse_number(raw) == expected


def test_met15_a_column_of_accounting_cells_profiles_as_numbers():
    csv = 'Month,Expenses\n2025-01," $ (1,234.00) "\n2025-02," $ -   "\n2025-03," $ 500.00 "\n'
    [ds] = parse_upload(csv.encode(), "t.csv")
    assert ds.columns[1].dtype == "number"
    assert [r["Expenses"] for r in ds.rows] == [-1234.0, 0.0, 500.0]


# ---------------------------------------------------------------- MET-16

MONTHLY_PL = "Month,Subscription revenue,Net income\n" + "".join(
    f"{p},10000,-5000\n" for p in months(2024, 1, 10)
)
QUARTERLY_CASH = "Quarter,Cash\nQ1 2024,135000\nQ2 2024,120000\nQ3 2024,105000\nQ4 2024,90000\n"


def test_met16_quarterly_target_drops_incomplete_quarters():
    ctx = ctx_of((MONTHLY_PL, "pl.csv"), (QUARTERLY_CASH, "cash.csv"), granularity="quarter")
    assert ctx.model.granularity == "quarter"
    assert ctx.figure("mrr").value == pytest.approx(10000)
    assert ctx.figure("burn").value == pytest.approx(5000)
    assert ctx.figure("runway_months").value == pytest.approx(18)
    assert any("Q4 2024 has 1 of 3 months" in n for n in ctx.model.notes)


def test_met16_monthly_target_says_the_quarterly_sheet_was_left_out():
    ctx = ctx_of((MONTHLY_PL, "pl.csv"), (QUARTERLY_CASH, "cash.csv"))
    assert ctx.model.granularity == "month"
    assert any("cash" in n and "quarter" in n for n in ctx.model.notes)


# ---------------------------------------------------------------- MET-17


def test_met17_rows_for_one_customer_in_a_period_add_up():
    csv = "Month,Customer,MRR\n2025-01,Acme,1000\n2025-01,Acme,500\n2025-01,Beta,1200\n"
    ctx = ctx_of((csv, "c.csv"))
    assert ctx.model.customer_mrr["2025-01"]["Acme"] == 1500
    assert ctx.figure("top_customer_share").value == pytest.approx(1500 / 2700)


def test_met17_rows_for_one_cohort_in_a_period_add_up():
    ds = DatasetInput(
        id="c",
        name="c",
        table_kind="cohort_mrr",
        granularity="month",
        transposed=False,
        rows=[
            {"period": "2025-01", "dim": "2024-01", "MRR": 100.0},
            {"period": "2025-01", "dim": "2024-01", "MRR": 50.0},
        ],
        mapping=[
            {"source_column": "Cohort", "canonical_field": "dim.cohort"},
            {"source_column": "MRR", "canonical_field": "value.mrr", "unit": "cad"},
        ],
    )
    ctx = build_context([ds], [], {"currency": "cad"})
    assert ctx.model.cohort_mrr["2024-01"]["2025-01"] == 150


# ---------------------------------------------------------------- MET-18


@pytest.mark.parametrize("unit", ["money", "$"])
def test_met18_money_units_are_the_reporting_currency(unit):
    ctx = build_context(
        [], [fact("mrr", 10_000, unit=unit), fact("round_size", 4_000_000, unit=unit)], {"currency": "cad"}
    )
    assert ctx.figure("mrr").unit == "cad" and ctx.figure("arr").value == 120_000
    assert ctx.figure("fact.round_size").unit == "cad"


def test_met18_unknown_fact_unit_is_missing_not_a_crash():
    ctx = build_context(
        [], [fact("cash", 10_000, unit="widgets"), fact("nrr", 1.1, unit="bananas")], {"currency": "cad"}
    )
    assert isinstance(ctx.figures["cash_balance"], MissingFigure)
    assert not isinstance(ctx.figures.get("nrr"), Figure)


# ---------------------------------------------------------------- MET-19


@pytest.mark.parametrize(
    "dates", [("2025-01-31", "2025-01-15", "2025-01-02"), ("31/01/2025", "15/01/2025", "02/01/2025")]
)
def test_met19_period_end_balance_is_the_latest_date(dates):
    csv = "Date,Cash,Revenue\n" + "".join(
        f"{d},{c},100\n" for d, c in zip(dates, (9000, 9500, 10000), strict=True)
    )
    [ds] = parse_upload(csv.encode(), "bank.csv")
    assert ds.rows[0]["Cash"] == 9000 and ds.rows[0]["Revenue"] == 300


# ---------------------------------------------------------------- MET-20


def test_met20_source_rows_are_file_lines():
    csv = "Acme P&L\nPrepared by finance\nMonth,Revenue\n2025-01,100\n2025-02,200\n2025-03,300\n"
    ctx = ctx_of((csv, "pl.csv"))
    _, ref, _ = ctx.model.source_ref("revenue.total", ["2025-03"])
    assert "row 6" in ref


def test_met20_merged_daily_rows_cite_every_line():
    csv = "Date,Revenue\n2025-01-02,100\n2025-01-15,100\n\n2025-02-01,100\n"
    ctx = ctx_of((csv, "pl.csv"))
    _, ref, _ = ctx.model.source_ref("revenue.total", ["2025-01"])
    assert "rows 2–3" in ref
    _, ref, _ = ctx.model.source_ref("revenue.total", ["2025-02"])
    assert "row 5" in ref


def test_met20_spreadsheet_rows_are_sheet_rows():
    import io

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A3"], ws["A5"], ws["B5"] = "Acme P&L", "Month", "Revenue"
    ws["A6"], ws["B6"], ws["A8"], ws["B8"] = "2025-01", 100, "2025-02", 200  # a blank row 7 in between
    buf = io.BytesIO()
    wb.save(buf)
    [ds] = parse_upload(buf.getvalue(), "pl.xlsx")
    assert [r["_src_rows"] for r in ds.rows] == [[6], [8]]


# ---------------------------------------------------------------- MET-23


@pytest.mark.parametrize(
    "text,scale",
    [
        ("Acme P&L (in $ thousands)", 1e3),
        ("All amounts in $ millions", 1e6),
        ("Revenue (k$)", 1e3),
        ("Revenue (m$)", 1e6),
        ("(in CAD thousands)", 1e3),
        ("Revenue ($000)", 1e3),
        ("In thousands of dollars", 1e3),
        ("Revenue", 1.0),
        ("Marketing spend", 1.0),
        ("Cost of goods sold", 1.0),
    ],
)
def test_met23_common_scale_headers(text, scale):
    assert scale_of(text) == scale


# ---------------------------------------------------------------- MET-24


@pytest.mark.parametrize(
    "name,bad",
    [
        ("Cost of sales - services", "revenue.one_time"),
        ("Services expense", "revenue.one_time"),
        ("Consulting fees", "revenue.one_time"),
        ("SaaS tools", "revenue.recurring"),
        ("Platform fees", "revenue.recurring"),
        ("Sales tax collected", "opex.s_and_m"),
        ("Cash collected", "cash_balance"),
    ],
)
def test_met24_expense_lines_never_map_to_revenue(name, bad):
    target, _ = _rule_match(name)
    assert target != bad
    assert not (target or "").startswith("revenue")


def test_met24_cost_and_tax_lines_get_a_sensible_target():
    assert _rule_match("Cost of sales - services")[0] == "cogs"
    assert _rule_match("Sales tax collected")[0] is None
    assert _rule_match("Cash collected")[0] is None
    assert _rule_match("Services expense")[0] is not None


@pytest.mark.parametrize(
    "name,target",
    [
        ("Subscription revenue", "revenue.recurring"),
        ("Professional services", "revenue.one_time"),
        ("Implementation fees", "revenue.one_time"),
        ("Subscription fees", "revenue.recurring"),
        ("Services revenue", "revenue.one_time"),
        ("Cash", "cash_balance"),
        ("SR&ED tax credit", "non_dilutive_income"),
        ("Software subscriptions", "opex.g_and_a"),
        ("Cost of revenue", "cogs"),
    ],
)
def test_met24_revenue_lines_still_map(name, target):
    assert _rule_match(name)[0] == target


# ---------------------------------------------------------------- MET-25


def test_met25_latest_fact_is_by_date_not_string_order():
    older, newer = fact("cash", 1, period="FY2024"), fact("cash", 2, period="2026-09")
    assert facts_by_metric([newer, older])["cash_balance"].value == 2
    assert facts_by_metric([older, newer])["cash_balance"].value == 2
    no_period = fact("cash", 3)
    assert facts_by_metric([newer, no_period])["cash_balance"].value == 2
    pending, confirmed = (
        fact("cash", 4, period="2026-09", status="pending"),
        fact("cash", 5, period="2026-09"),
    )
    assert facts_by_metric([confirmed, pending])["cash_balance"].value == 5


# ---------------------------------------------------------------- MET-26


def test_met26_missing_rate_is_the_stated_reason():
    ds = DatasetInput(
        id="pl",
        name="pl",
        table_kind="series",
        granularity="month",
        transposed=False,
        rows=[{"period": "2025-01", "rev": 1000.0}],
        mapping=[{"source_column": "rev", "canonical_field": "revenue.total", "unit": "usd"}],
    )
    ctx = build_context([ds], [], {"currency": "cad"}, fx=None)
    rev = ctx.figures["revenue"]
    assert isinstance(rev, MissingFigure) and "No USD→CAD rate" in rev.reason and "2025-01" in rev.reason


# ---------------------------------------------------------------- MET-28


def test_met28_yearly_labels():
    csv = "Year,Subscription revenue\n2022,100000\n2023,150000\n2024,200000\n"
    ctx = ctx_of((csv, "y.csv"))
    assert ctx.model.granularity == "year"
    assert "year over year" in ctx.figure("revenue_growth").label
    assert isinstance(ctx.figures["cmgr_recent"], MissingFigure)


# ---------------------------------------------------------------- MET-29


def test_met29_derived_figures_keep_file_provenance():
    ctx = sample_context("harbor-labs")
    for fid in ("burn", "runway_months"):
        assert ctx.figure(fid).source.dataset_id, fid
    ds, ref, _ = ctx.model.source_ref("net_income", ["2026-Q3"])
    assert ds and ref.startswith("derived: ") and "financials_quarterly" in ref
