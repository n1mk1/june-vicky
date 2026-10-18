"""The pitch against the founder's own knowledge base: PowerPoint decks, the statement checker
(vicky.pitch.verify), notes interrupts and the discrepancy questions after the pitch."""

import io

import pytest
from pptx import Presentation
from pptx.util import Inches

from vicky.areopagus import registry
from vicky.court import minds as court_minds
from vicky.court import session as court
from vicky.model.context import sample_context
from vicky.pitch import session as pitch
from vicky.pitch import verify


def _pptx() -> bytes:
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Traction"
    s1.placeholders[1].text = "Three hospital pilots live\nLaunched in Germany in March"
    s1.notes_slide.notes_text_frame.text = "say this slowly"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Team"
    t = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    t.cell(0, 0).text, t.cell(0, 1).text = "Role", "People"
    t.cell(1, 0).text, t.cell(1, 1).text = "Engineers", "12"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _passage(text: str, path: str = "company/Pilots.md") -> dict:
    chunk = {
        "id": "c1",
        "note_id": "n1",
        "path": path,
        "title": "Pilots",
        "heading": "",
        "text": text,
        "cite": f"[kb:{path}]",
    }
    return {"path": path, "cite": chunk["cite"], "heading": "", "text": text, "chunk": chunk}


def test_pptx_slides_are_read_as_text_without_speaker_notes():
    slides = pitch.read_deck(_pptx(), "deck.pptx")
    assert [s["n"] for s in slides] == [1, 2]
    assert slides[0]["title"] == "Traction" and slides[0]["image"] is None
    assert "Three hospital pilots live" in slides[0]["bullets"] and "say this slowly" not in slides[0]["text"]
    assert "Engineers · 12" in slides[1]["bullets"]


def test_deck_reader_names_both_formats_when_refusing():
    with pytest.raises(pitch.PitchError, match="PowerPoint"):
        pitch.read_deck(b"GIF89a", "deck.gif")


def test_offline_assertions_keep_checkable_sentences_only():
    found = verify.offline_assertions(
        "Hello everyone. We have three hospital pilots live in Ontario. Thanks!"
    )
    assert [a["text"] for a in found] == ["We have three hospital pilots live in Ontario."]


def test_offline_rating_supported_new_and_unclear():
    a = {"text": "We signed 3 hospital pilots", "topic": "pilots"}
    ok = verify.offline_rating(a, [_passage("By March we had signed 3 hospital pilots in Ontario.")], True)
    assert (
        ok["verdict"] == "supported" and ok["quote"] == "By March we had signed 3 hospital pilots in Ontario."
    )
    assert ok["cite"] == "[kb:company/Pilots.md]" and ok["chunk"]["note_id"] == "n1"
    assert verify.offline_rating(a, [], False)["verdict"] == "new"
    assert verify.offline_rating(a, [_passage("We talk to hospitals.")], True)["verdict"] == "unclear"


def test_a_model_verdict_must_quote_its_passage_and_adds_no_number():
    a = {"text": "We have 5 pilots", "topic": "pilots"}
    ps = [_passage("We have 3 pilots running.")]
    bad = verify.Rating(n=1, verdict="contradicted", passage=1, quote="We have 4 pilots", why="Notes say 4.")
    assert verify.clean_rating(bad, a, ps)["verdict"] == "unclear"
    good = verify.Rating(
        n=1,
        verdict="contradicted",
        passage=1,
        quote="We have 3 pilots running.",
        why="The notes count 3 pilots, not 5. Revenue is 900k.",
    )
    r = verify.clean_rating(good, a, ps)
    assert r["verdict"] == "contradicted" and r["cite"] == "[kb:company/Pilots.md]"
    assert "900" not in r["why"] and "3 pilots" in r["why"]


async def test_check_without_an_index_returns_nothing(monkeypatch):
    async def none(_t):
        return [], False, False

    monkeypatch.setattr(verify, "_retrieve", none)
    assert await verify.check("We have three hospital pilots live in Ontario.") == []


async def test_check_offline_with_passages(monkeypatch):
    async def found(t):
        if "pilots" in t:
            return [_passage("We run 3 hospital pilots in Ontario today.")], True, True
        return [], False, True

    monkeypatch.setattr(verify, "_retrieve", found)
    monkeypatch.setattr(verify.llm, "enabled", lambda: False)
    rows = await verify.check(
        "We run 3 hospital pilots in Ontario. We are opening an office in Lisbon next year."
    )
    by = {r["text"]: r["verdict"] for r in rows}
    assert by["We run 3 hospital pilots in Ontario."] == "supported"
    assert by["We are opening an office in Lisbon next year."] == "new"


def test_a_statement_the_notes_contradict_interrupts_with_the_note_quoted():
    ctx = sample_context("tidewell")
    panel = [registry.builtins()[i] for i in ("aristotle", "thales")]
    k = {
        "text": "We have five hospital pilots",
        "verdict": "contradicted",
        "quote": "Two hospital pilots are live.",
        "cite": "[kb:company/Pilots.md]",
    }
    p, chosen, score = pitch._candidate(panel, [], [k])
    assert chosen is k and score >= pitch.INTERRUPT_BAR
    notes = {
        "company/Pilots.md": {
            "id": "n1",
            "title": "Pilots",
            "content": "Status. Two hospital pilots are live.",
        }
    }
    line = pitch._kb_line(p, k, "We have five hospital pilots.", ctx, notes)
    assert '"Two hospital pilots are live."' in line and "[kb:company/Pilots.md]" in line
    # Without the note to check the quote against, the line names no quote.
    assert "Two hospital" not in pitch._kb_line(p, k, "We have five hospital pilots.", ctx)
    assert pitch._candidate(panel, [], [{**k, "verdict": "new"}]) is None  # new claims wait for the Q&A


def test_discrepancies_put_contradictions_first_and_dedupe():
    kb = [
        {"text": "We launched in Spain", "verdict": "new"},
        {
            "text": "We have 5 pilots",
            "verdict": "contradicted",
            "quote": "We have 3 pilots.",
            "cite": "[kb:a.md]",
        },
        {
            "text": "We have 5 pilots.",
            "verdict": "contradicted",
            "quote": "We have 3 pilots.",
            "cite": "[kb:a.md]",
        },
        {"text": "Our churn is low", "verdict": "supported"},
    ]
    out = pitch.discrepancies(kb)
    assert [k["text"] for k in out] == ["We have 5 pilots", "We launched in Spain"]


def test_discrepancies_become_doubts_with_code_written_questions():
    ctx = sample_context("tidewell")
    panel = [registry.builtins()[i] for i in ("aristotle", "thales")]
    state = {"vcs": {p.id: {"hyps": []} for p in panel}}
    chunk = _passage("We have 3 pilots.")["chunk"]
    court.seat_discrepancies(
        state,
        panel,
        [
            {
                "id": "k1",
                "text": "We have 5 pilots",
                "verdict": "contradicted",
                "quote": "We have 3 pilots.",
                "cite": "[kb:company/Pilots.md]",
                "chunk": chunk,
            },
            {"id": "k2", "text": "We launched in Spain", "verdict": "new"},
        ],
    )
    hyps = sorted((h for v in state["vcs"].values() for h in v["hyps"]), key=lambda h: h["attack_id"])
    assert len(hyps) == 2 and hyps[0]["kb"] == [chunk] and hyps[1]["kb"] == []
    p = next(p for p in panel if hyps[0] in state["vcs"][p.id]["hyps"])
    q = court_minds.question_for(p, hyps[0], ctx, set())
    assert '"We have 3 pilots."' in q and "Which is right" in q
    q2 = court_minds.question_for(p, hyps[1], ctx, set())
    assert "nothing in your notes covers it" in q2
