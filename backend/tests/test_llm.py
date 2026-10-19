"""The LLM path on recorded responses: client, minds, mapper and FR-7 extraction. No live calls."""

import asyncio
import io
import json

import pytest
from fastapi import HTTPException

from tests.llm_fakes import (
    api_error,
    json_reply,
    make_pdf,
    message,
    queue,
    text_reply,
    tool_use_reply,
)
from vik import auth, llm
from vik.areopagus import builder, registry
from vik.audit import audit_text
from vik.config import get_settings
from vik.court import minds as offline
from vik.court import referee
from vik.llm import docs, mapper
from vik.llm import minds as agents
from vik.llm.client import (
    INVESTOR_ROUTES,
    ROUTES,
    VIK_ROUTES,
    BudgetExceeded,
    LLMError,
    LLMRefusal,
    LLMUnavailable,
    get_llm,
    model_for,
    over_budget,
    usage_usd,
    with_timeout,
)
from vik.llm.schemas import ChairDecision, ClaimList, DocFact, DocFactList, Memo, MemoHypothesis
from vik.model.context import sample_context
from vik.model.loader import parse_upload

INVENTED = "Your ARR is $9.9M, so why is growth stalling?"


@pytest.fixture(scope="module")
def ctx():
    return sample_context("tidewell")


@pytest.fixture
def james():
    return registry.builtins()["aristotle"]


def by_task(**handlers):
    """Responder that dispatches on the agent prompt a call carries; each value is a list (queue) or a callable."""
    pending = {k: (list(v) if isinstance(v, list) else v) for k, v in handlers.items()}

    def respond(call):
        h = pending.get(call["task"])
        if h is None:
            raise AssertionError(f"unexpected {call['task']} call")
        if isinstance(h, list):
            if not h:
                raise AssertionError(f"too many {call['task']} calls")
            return h.pop(0)
        return h(call)

    return respond


# ---------------------------------------------------------------- client


async def test_parse_sends_route_settings_and_records_spend(live_llm):
    rec = live_llm.install(queue(json_reply({"claims": []}, model="gemini-3.5-flash-lite")))
    out = await get_llm().parse(
        "claims", ClaimList, system="Task: list claims", messages=[{"role": "user", "content": "hi"}]
    )
    assert out.claims == []
    body = rec.calls[0]["body"]
    assert body["model"] == "gemini-3.5-flash-lite" and body["thinking_level"] == "low"
    assert body["system"] == "Task: list claims" and body["max_tokens"] == ROUTES["claims"].max_tokens
    assert body["response_schema"]["type"] == "object" and "claims" in body["response_schema"]["properties"]
    assert "$defs" not in json.dumps(body["response_schema"])  # references are inlined
    assert rec.calls[0]["raw_body"]["generationConfig"]["responseMimeType"] == "application/json"
    assert live_llm.spend == pytest.approx((1000 * 0.30 + 200 * 2.50) / 1e6)


def test_the_default_models_are_all_on_the_free_tier_and_spread_over_three_quotas():
    for name in VIK_ROUTES:
        assert (model_for(ROUTES[name]), ROUTES[name].effort) == ("gemini-3.8-flash", "low"), name
    for name in INVESTOR_ROUTES:
        assert (model_for(ROUTES[name]), ROUTES[name].effort) == ("gemini-3.6-flash", "low"), name
    assert (
        model_for(ROUTES["claims"]) == "gemini-3.5-flash-lite"
        and model_for(ROUTES["doc_facts"]) == "gemini-3.8-flash"
    )
    assert get_settings().gemini_tier == "free"
    assert "gemini-4-argon" not in {model_for(r) for r in ROUTES.values()}  # not on the free tier


async def test_vik_chairs_at_a_low_thinking_level(live_llm):
    decision = {
        "action": "call_on",
        "vc_id": "a",
        "intent": "new_topic",
        "hypothesis_id": "a.h1",
        "note": None,
    }
    rec = live_llm.install(queue(json_reply(decision)))
    await get_llm().parse("chair", ChairDecision, system="chair", messages=[{"role": "user", "content": "x"}])
    call = rec.calls[0]
    # Thinking can't be turned off on these models: the thinking level is the only control.
    assert call["body"]["model"] == "gemini-3.8-flash" and call["body"]["thinking_level"] == "low"
    assert call["headers"].get("x-goog-api-key") == "test-key"


async def test_investor_lines_go_to_flash_with_the_investor_role(live_llm):
    rec = live_llm.install(queue(text_reply("Which months is that?")))
    await get_llm().text("speak", system="speak", messages=[{"role": "user", "content": "x"}])
    body = rec.calls[0]["body"]
    assert body["model"] == "gemini-3.6-flash" and body["thinking_level"] == "low"
    assert "Your role" in agents.INVESTOR_ROLE and "foresight" in agents.INVESTOR_ROLE
    assert "Vik" in agents.VIK_ROLE


async def test_refusal_is_checked_before_content(live_llm):
    live_llm.install(
        queue(
            message([], stop_reason="SAFETY"),
            message([], blocked="PROHIBITED_CONTENT"),
        )
    )
    with pytest.raises(LLMRefusal) as e:
        await get_llm().parse("claims", ClaimList, system="s", messages=[{"role": "user", "content": "x"}])
    assert e.value.category == "safety"
    assert live_llm.spend > 0  # the declined call is still accounted
    with pytest.raises(LLMRefusal) as e:
        await get_llm().parse("claims", ClaimList, system="s", messages=[{"role": "user", "content": "x"}])
    assert e.value.category == "PROHIBITED_CONTENT"  # a blocked prompt has no candidate at all


async def test_truncated_output_and_server_errors_are_llm_errors(live_llm):
    live_llm.install(
        queue(
            message([{"text": '{"claims": ['}], stop_reason="max_tokens"),
            api_error(503, "busy"),
            api_error(400, "bad schema"),
            api_error(400, "API key not valid. Please pass a valid API key."),
            api_error(429, "quota"),
        )
    )
    with pytest.raises(LLMError):
        await get_llm().parse("claims", ClaimList, system="s", messages=[{"role": "user", "content": "x"}])
    with pytest.raises(LLMUnavailable):
        await get_llm().parse("claims", ClaimList, system="s", messages=[{"role": "user", "content": "x"}])
    with pytest.raises(LLMError, match="bad request") as e:
        await get_llm().parse("claims", ClaimList, system="s", messages=[{"role": "user", "content": "x"}])
    assert not isinstance(e.value, LLMUnavailable)
    with pytest.raises(LLMUnavailable, match="credentials rejected"):  # how Gemini reports a bad key
        await get_llm().parse("claims", ClaimList, system="s", messages=[{"role": "user", "content": "x"}])
    with pytest.raises(LLMUnavailable, match="rate limited"):
        await get_llm().parse("claims", ClaimList, system="s", messages=[{"role": "user", "content": "x"}])


# ---------------------------------------------------------------- free tier


def _free(monkeypatch, rpm: int = 10) -> None:
    monkeypatch.setenv("GEMINI_TIER", "free")
    monkeypatch.setenv("GEMINI_RPM", str(rpm))
    get_settings.cache_clear()


async def test_a_free_key_costs_nothing_and_no_dollar_budget_blocks_it(live_llm, monkeypatch):
    _free(monkeypatch)
    rec = live_llm.install(lambda _call: json_reply({"claims": []}))
    live_llm.spend = get_settings().daily_usd_budget + 1  # a paid day's spend left over
    await get_llm().parse("claims", ClaimList, system="s", messages=[{"role": "user", "content": "x"}])
    assert len(rec.calls) == 1 and live_llm.spend == get_settings().daily_usd_budget + 1  # nothing added
    assert not await over_budget()
    await auth.require_budget()  # LLM routes stay open (no 402)


async def test_the_pacer_keeps_each_model_under_its_requests_per_minute(live_llm, monkeypatch):
    from vik.llm.client import PACER

    _free(monkeypatch, rpm=2)
    await PACER.acquire("gemini-3.8-flash")
    await PACER.acquire("gemini-3.8-flash")
    with pytest.raises(TimeoutError):  # the third request waits for a slot
        await asyncio.wait_for(PACER.acquire("gemini-3.8-flash"), 0.2)
    await asyncio.wait_for(
        PACER.acquire("gemini-3.5-flash-lite"), 0.2
    )  # another model has its own quota (1 of 2)
    rec = live_llm.install(lambda _call: json_reply({"claims": []}))
    await get_llm().parse("claims", ClaimList, system="s", messages=[{"role": "user", "content": "x"}])
    with pytest.raises(
        LLMUnavailable, match="no answer within|timed out"
    ):  # its route timeout bounds the wait
        await with_timeout(
            get_llm().parse("claims", ClaimList, system="s", messages=[{"role": "user", "content": "x"}]), 0.2
        )
    assert len(rec.calls) == 1


async def test_a_429_rests_that_model_and_honours_retry_delay(live_llm, monkeypatch):
    from vik.llm.client import PACER

    _free(monkeypatch)
    quota = (
        429,
        {
            "error": {
                "code": 429,
                "message": "Quota exceeded",
                "status": "RESOURCE_EXHAUSTED",
                "details": [{"@type": "type.googleapis.com/google.rpc.RetryInfo", "retryDelay": "37s"}],
            }
        },
    )
    rec = live_llm.install(
        lambda call: quota if call["body"]["model"] == "gemini-3.5-flash-lite" else text_reply("ok")
    )
    with pytest.raises(LLMUnavailable, match="rate limited"):
        await get_llm().parse("claims", ClaimList, system="s", messages=[{"role": "user", "content": "x"}])
    import time

    assert 30 < PACER.rest_until["gemini-3.5-flash-lite"] - time.monotonic() <= 37
    with pytest.raises(LLMUnavailable, match="rests for"):  # no request is sent while it rests
        await get_llm().parse("claims", ClaimList, system="s", messages=[{"role": "user", "content": "x"}])
    assert len(rec.calls) == 1
    assert await get_llm().text("speak", system="s", messages=[{"role": "user", "content": "x"}]) == "ok"
    assert await agents.extract_claims("We have 40 customers.") == offline_claims("We have 40 customers.")


async def test_budget_blocks_calls_and_routes(live_llm, monkeypatch):
    rec = live_llm.install(queue())
    live_llm.spend = get_settings().daily_usd_budget + 1
    with pytest.raises(BudgetExceeded):
        await get_llm().text("speak", system="s", messages=[{"role": "user", "content": "x"}])
    assert rec.calls == []
    with pytest.raises(HTTPException) as e:
        await auth.require_budget()
    assert e.value.status_code == 402
    monkeypatch.setenv("LLM_MODE", "offline")
    get_settings.cache_clear()
    await auth.require_budget()  # the offline path costs nothing and is never blocked


def test_enabled_needs_a_key_and_respects_offline(monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "")  # beats a key in a local .env
    monkeypatch.setenv("GOOGLE_API_KEY", "")
    monkeypatch.setenv("LLM_MODE", "auto")
    get_settings.cache_clear()
    assert llm.enabled() is False
    monkeypatch.setenv("GEMINI_API_KEY", "k")
    get_settings.cache_clear()
    assert llm.enabled() is True
    monkeypatch.setenv("GEMINI_API_KEY", "")
    monkeypatch.setenv("GOOGLE_API_KEY", "k")  # the Google SDKs' name works too
    get_settings.cache_clear()
    assert llm.enabled() is True
    monkeypatch.setenv("LLM_MODE", "offline")
    get_settings.cache_clear()
    assert llm.enabled() is False
    monkeypatch.undo()
    get_settings.cache_clear()


def test_usage_is_priced_at_list_rates(monkeypatch):
    from types import SimpleNamespace

    monkeypatch.setenv("GEMINI_TIER", "paid")
    get_settings.cache_clear()

    # The prompt count includes the cached tokens; thinking is billed as output.
    u = SimpleNamespace(
        prompt_token_count=2_000_000,
        cached_content_token_count=1_000_000,
        candidates_token_count=500_000,
        thoughts_token_count=500_000,
    )
    assert usage_usd("gemini-4-argon", u) == pytest.approx(4 + 20 + 0.2)
    assert usage_usd("gemini-3.8-flash", u) == pytest.approx(1.5 + 7.5 + 0.15)
    assert usage_usd("gemini-3.8-flash-001", u) == usage_usd("gemini-3.8-flash", u)  # a pinned version
    assert usage_usd("some-new-model", u) >= usage_usd("gemini-4-argon", u)  # unknown → dearest
    monkeypatch.setenv("GEMINI_TIER", "free")
    get_settings.cache_clear()
    assert usage_usd("gemini-3.8-flash", u) == 0  # a free key costs nothing
    monkeypatch.undo()
    get_settings.cache_clear()


# ---------------------------------------------------------------- minds: offline passthrough


async def test_without_a_key_the_minds_are_the_offline_minds(ctx, james, monkeypatch):
    monkeypatch.setenv("LLM_MODE", "offline")
    get_settings.cache_clear()
    try:
        assert await agents.build_memos([james], ctx) == offline.build_memos([james], ctx)
        hyp = offline.build_memos([james], ctx)["aristotle"]["hypotheses"][0]
        spoken = await agents.speak(james, hyp, ctx, "new_topic", set())
        assert spoken.text == offline.question_for(james, hyp, ctx, set()) and spoken.source == "offline"
        answer = "Our ARR is $1.29M and churn is about 1%."
        assert await agents.extract_claims(answer) == offline_claims(answer)
    finally:
        monkeypatch.undo()
        get_settings.cache_clear()


def offline_claims(answer):
    from vik.model.truth import quick_claims

    return quick_claims(answer)


# ---------------------------------------------------------------- memos


async def test_memo_uses_kit_tools_and_code_keeps_only_anchored_hypotheses(live_llm, ctx, james):
    base = offline.build_memos([james], ctx)["aristotle"]
    owned = base["hypotheses"][0]
    real_fig = next(f.id for f in ctx.available() if f.id not in owned["evidence"])
    memo = {
        "thesis": "Growth is slowing and the plan assumes it won't.",
        "hypotheses": [
            {
                "attack_id": owned["attack_id"],
                "claim": "Growth is decelerating",
                "lens": owned["lens"],
                "evidence": owned["evidence"],
                "severity": 0.9,
                "confidence": 0.99,
            },
            {
                "attack_id": None,
                "claim": "Hidden doubt with no data",
                "lens": "market",
                "evidence": ["not_a_figure"],
                "severity": 0.9,
                "confidence": 0.9,
            },
            {
                "attack_id": None,
                "claim": "The plan leans on one line",
                "lens": "forecast",
                "evidence": [real_fig],
                "severity": 0.99,
                "confidence": 0.95,
            },
        ],
        "mind_changers": ["Cohort data by month"],
        "opening_question": INVENTED,
    }
    rec = live_llm.install(
        by_task(
            memo=[
                tool_use_reply("get_metric", {"id": "arr"}, model="gemini-4-argon"),
                json_reply(memo, model="gemini-4-argon"),
            ]
        )
    )
    out = (await agents.build_memos([james], ctx))["aristotle"]
    assert out["source"] == "llm"
    assert [h["id"] for h in out["hypotheses"]] == ["aristotle.h1", "aristotle.h2"]
    first, new = out["hypotheses"]
    assert first["attack_id"] == owned["attack_id"] and first["confidence"] == 0.9  # clamped
    assert new["evidence"] == [real_fig] and new["severity"] == 0.7  # a model-only doubt can't outrank code's
    assert (
        out["opening_question"] != INVENTED
        and audit_text(out["opening_question"], ctx.available(), first.get("quotes")).passed
    )
    # the kit tool ran in code and its Figure went back to the model
    second = rec.calls[1]["body"]["messages"]
    result = next(
        b
        for m in second
        if m["role"] == "user" and isinstance(m["content"], list)
        for b in m["content"]
        if b.get("type") == "tool_result"
    )
    assert json.loads(result["content"])["id"] == "arr"
    # the investor researches on Gemini 3.6 Flash (its own free quota), in role; its framing of an attack-list doubt is kept as its angle
    assert rec.calls[0]["body"]["model"] == "gemini-3.6-flash"
    assert "Your role" in rec.calls[0]["body"]["system"]
    assert "get_metric" in {t["name"] for t in rec.calls[0]["body"]["tools"]}
    # the model's turn went back as it came, thought signature included
    assert rec.calls[1]["raw_body"]["contents"][1]["parts"][0]["thoughtSignature"] == "c2ln"
    assert first["angle"] == "Growth is decelerating" and "angle" not in new


async def test_an_angle_with_a_number_is_dropped_and_a_kept_one_reaches_the_line(live_llm, ctx, james):
    base = offline.build_memos([james], ctx)["aristotle"]
    owned = base["hypotheses"][0]
    m = Memo(
        thesis="x",
        hypotheses=[
            MemoHypothesis(
                attack_id=owned["attack_id"],
                claim="Growth fell 40% last quarter",
                lens=owned["lens"],
                evidence=owned["evidence"],
                severity=0.5,
                confidence=0.5,
            )
        ],
        mind_changers=[],
        opening_question="",
    )
    assert "angle" not in agents.memo_from_llm(m, james, ctx, base, set())["hypotheses"][0]
    rec = live_llm.install(by_task(speak=[text_reply("What do you do if the slowdown is the new normal?")]))
    await agents.speak(james, {**owned, "angle": "The plan ignores the slowdown"}, ctx, "new_topic", set())
    assert (
        "Your angle, from your research: The plan ignores the slowdown"
        in rec.calls[0]["body"]["messages"][0]["content"]
    )


async def test_memo_failure_falls_back_per_investor(live_llm, ctx, james):
    live_llm.install(by_task(memo=[api_error(500, "x")]))
    assert (await agents.build_memos([james], ctx))["aristotle"] == offline.build_memos([james], ctx)[
        "aristotle"
    ]


# ---------------------------------------------------------------- speaking


async def test_a_line_that_fails_the_audit_is_regenerated_once(live_llm, ctx, james):
    hyp = offline.build_memos([james], ctx)["aristotle"]["hypotheses"][0]
    good = "Walk me through which months that growth rate covers and how you calculate it."
    rec = live_llm.install(by_task(speak=[text_reply(INVENTED), text_reply(good)]))
    spoken = await agents.speak(james, hyp, ctx, "new_topic", set())
    assert (spoken.text, spoken.regenerated, spoken.source) == (good, True, "llm")
    retry = rec.calls[1]["body"]["messages"]
    assert retry[-1]["role"] == "user" and "failed the fact audit" in retry[-1]["content"]


async def test_two_failed_lines_fall_back_to_the_offline_line(live_llm, ctx, james):
    hyp = offline.build_memos([james], ctx)["aristotle"]["hypotheses"][0]
    live_llm.install(by_task(speak=[text_reply(INVENTED), text_reply("We value you at $20M pre-money.")]))
    spoken = await agents.speak(james, hyp, ctx, "new_topic", set())
    assert spoken.source == "offline" and spoken.regenerated is True
    assert spoken.text == offline.question_for(james, hyp, ctx, set())


# ---------------------------------------------------------------- chair


def _ranked():
    return referee.rank(
        [
            referee.Bid("aristotle", 0.8, "new_topic", "aristotle.h1", turns_since_spoke=5),
            referee.Bid("thales", 0.6, "new_topic", "thales.h1", turns_since_spoke=5),
        ],
        answer_weak=False,
    )


def _state():
    return {
        "vcs": {
            v: {
                "hyps": [
                    {
                        "id": f"{v}.h1",
                        "claim": "doubt",
                        "lens": "forecast",
                        "confidence": 0.7,
                        "asked": 0,
                        "status": "open",
                    }
                ]
            }
            for v in ("aristotle", "thales")
        }
    }


async def test_chair_can_pick_any_ranked_candidate(live_llm):
    panel = {p: registry.builtins()[p] for p in ("aristotle", "thales")}
    live_llm.install(
        by_task(
            chair=[
                json_reply(
                    {
                        "action": "call_on",
                        "vc_id": "thales",
                        "intent": "new_topic",
                        "hypothesis_id": "thales.h9",
                        "note": "fresh lens",
                    }
                ),
                json_reply(
                    {
                        "action": "call_on",
                        "vc_id": "thales",
                        "intent": "follow_up",
                        "hypothesis_id": "thales.h1",
                        "note": None,
                    }
                ),
            ]
        )
    )
    pick = await agents.chair(_ranked(), _state(), panel)
    assert pick.candidate.vc_id == "thales" and not pick.fallback and pick.valid
    assert pick.decision["hypothesis_id"] == "thales.h1"  # an unknown hypothesis becomes the candidate's own
    pick = await agents.chair(_ranked(), _state(), panel)
    assert (
        pick.fallback and pick.candidate.vc_id == "aristotle"
    )  # a follow-up the referee didn't allow is invalid


async def test_invalid_or_late_chair_falls_back_to_the_top_candidate(live_llm, monkeypatch):
    panel = {p: registry.builtins()[p] for p in ("aristotle", "thales")}
    live_llm.install(
        by_task(
            chair=[
                json_reply(
                    {
                        "action": "call_on",
                        "vc_id": "zed",
                        "intent": "new_topic",
                        "hypothesis_id": None,
                        "note": None,
                    }
                )
            ]
        )
    )
    pick = await agents.chair(_ranked(), _state(), panel)
    assert pick.fallback and pick.candidate.vc_id == "aristotle" and "not an eligible" in pick.reason

    async def slow(_call):
        await asyncio.sleep(0.5)
        return json_reply(
            {
                "action": "call_on",
                "vc_id": "thales",
                "intent": "new_topic",
                "hypothesis_id": "thales.h1",
                "note": None,
            }
        )

    monkeypatch.setenv("CHAIR_TIMEOUT_S", "0.05")
    get_settings.cache_clear()
    live_llm.install(slow)
    pick = await agents.chair(_ranked(), _state(), panel)
    assert pick.fallback and pick.candidate.vc_id == "aristotle"


# ---------------------------------------------------------------- beliefs


def _vc_state():
    return {
        "hyps": [
            {
                "id": "aristotle.h1",
                "attack_id": "growth_decel",
                "claim": "Growth is slowing",
                "lens": "forecast",
                "evidence": ["cmgr_12m"],
                "confidence": 0.7,
                "severity": 0.8,
                "status": "open",
                "asked": 1,
                "weak_answers": 0,
                "quotes": [],
            }
        ],
        "last_spoke": 0,
    }


async def test_no_concession_without_a_matching_figure_holds_on_the_llm_path(live_llm, ctx, james):
    drop = {
        "deltas": [{"hyp_id": "aristotle.h1", "new_confidence": 0.1, "why": "They sounded confident"}],
        "bid": 0.4,
        "intent": "follow_up",
        "target_hyp": "aristotle.h1",
    }
    live_llm.install(by_task(bid=[json_reply(drop), json_reply(drop)]))
    state = _vc_state()
    out = await agents.update_and_bid(
        james, state, [], False, "thales", "thales.h1", 2, set(), ctx=ctx, answer="Trust me."
    )
    assert state["hyps"][0]["confidence"] == 0.7 and out["deltas"] == []
    assert out["intent"] != "follow_up"  # not this investor's thread
    match = [{"status": "matches", "figure_id": "cmgr_12m", "text": "growth"}]
    out = await agents.update_and_bid(
        james, state, match, False, "aristotle", "aristotle.h1", 0, set(), ctx=ctx
    )
    assert state["hyps"][0]["confidence"] == pytest.approx(0.35)  # accepted, but at most MAX_DELTA per answer
    assert out["deltas"][0]["hyp_id"] == "aristotle.h1"


# ---------------------------------------------------------------- judge and claims


async def test_judge_scores_are_clamped_and_rewrites_audited(live_llm, ctx, james):
    hyp = offline.build_memos([james], ctx)["aristotle"]["hypotheses"][0]
    answer = "We grew fast last year and we're fixing the pipeline."
    js = {
        "evidence": 7,
        "clarity": 3.04,
        "composure": -1,
        "investor_fit": 4,
        "good": "Owned the slowdown.",
        "missing": "The growth figure and its window.",
        "strong_answer": "Our ARR is $9.9M and growing 40% a month.",
    }
    live_llm.install(by_task(judge=[json_reply(js, model="gemini-4-argon")]))
    base = offline.judge(answer, hyp, [], ctx, "q")
    out = await agents.judge(answer, hyp, [], ctx, "q")
    assert out["dims"] == {"evidence": 5.0, "clarity": 3.0, "composure": 0.0, "investor_fit": 4.0}
    assert out["rationale"]["good"] == "Owned the slowdown."
    assert (
        out["rationale"]["strong_answer"] == base["rationale"]["strong_answer"]
    )  # invented numbers rejected
    assert (
        out["accuracy"] == base["accuracy"]
        and out["rationale"]["correct_figure"] == base["rationale"]["correct_figure"]
    )


async def test_claims_must_quote_the_founder_with_the_number(live_llm):
    answer = "Honestly our ARR is $1.29M. We sign about 14 new logos a month."
    claims = {
        "claims": [
            {
                "text": "our ARR is $1.29M",
                "metric": "arr",
                "value": 1290000,
                "unit": "money",
                "period": None,
                "is_forecast": False,
            },
            {
                "text": "NRR is 130%",
                "metric": "nrr",
                "value": 130,
                "unit": "pct",
                "period": None,
                "is_forecast": False,
            },
            {
                "text": "We sign about 14 new logos",
                "metric": "new customers",
                "value": 40,
                "unit": "count",
                "period": None,
                "is_forecast": False,
            },
            {
                "text": "We sign about 14 new logos a month",
                "metric": "new customers",
                "value": 14,
                "unit": "count",
                "period": "month",
                "is_forecast": False,
            },
        ]
    }
    live_llm.install(by_task(claims=[json_reply(claims, model="gemini-3.5-flash-lite")]))
    out = await agents.extract_claims(answer)
    assert out[: len(offline_claims(answer))] == offline_claims(answer)  # code's own extraction comes first
    assert [c["metric"] for c in out].count("arr") == 1  # no duplicate of what code found
    assert "nrr" not in [c["metric"] for c in out]  # not in the answer
    added = [c for c in out if c["metric"] == "new customers"]
    assert [c["value"] for c in added] == [14]  # the 40 isn't in its quote


# ---------------------------------------------------------------- column mapper

CSV = (
    b"Month,Total Revenue,Subscription revenue,Hosting,Mystery Col\n"
    b"2025-01,100,90,10,5\n2025-02,110,99,11,6\n2025-03,120,108,12,7\n"
)


async def test_mapper_merges_gemini_with_the_heuristic(live_llm):
    ds = parse_upload(CSV, "numbers.csv")[0]
    proposal = {
        "columns": [
            {
                "source_column": "Mystery Col",
                "canonical_field": "opex.g_and_a",
                "unit": "usd",
                "confidence": 0.8,
                "note": "office costs",
            },
            {
                "source_column": "Subscription revenue",
                "canonical_field": "revenue.total",
                "unit": "",
                "confidence": 0.9,
                "note": "x",
            },
            {
                "source_column": "Hosting",
                "canonical_field": "made.up",
                "unit": "",
                "confidence": 0.99,
                "note": "x",
            },
            {
                "source_column": "Not a column",
                "canonical_field": "cogs",
                "unit": "",
                "confidence": 0.99,
                "note": "x",
            },
        ]
    }
    rec = live_llm.install(by_task(mapper=[json_reply(proposal)]))
    out, proposer = await mapper.propose(ds, "cad")
    by = {m.source_column: m for m in out}
    assert proposer == "gemini" and rec.calls[0]["body"]["model"] == "gemini-3.8-flash"
    assert (
        by["Mystery Col"].canonical_field == "opex.g_and_a" and by["Mystery Col"].unit == "cad"
    )  # unit set by code
    assert by["Subscription revenue"].canonical_field != "revenue.total"  # the total column keeps its field
    assert (
        by["Hosting"].canonical_field != "made.up" and by["Total Revenue"].canonical_field == "revenue.total"
    )


async def test_mapper_without_a_key_is_the_heuristic():
    from vik.model.mapping import propose_heuristic

    ds = parse_upload(CSV, "numbers.csv")[0]
    out, proposer = await mapper.propose(ds, "cad")
    assert proposer == "heuristic" and out == propose_heuristic(ds, "cad")


# ---------------------------------------------------------------- FR-13 Areopagus builder

DRAFT = {
    "name": "Rhea",
    "firm": "Seed fund",
    "archetype": "NRR obsessed SaaS partner",
    "cheque_role": "leads_round",
    "bio": "Former operator who backs SaaS companies with sticky revenue.",
    "priorities": [
        {"lens": "retention", "weight": 1.7},
        {"lens": "runway", "weight": 0.8},
        {"lens": "team", "weight": -0.3},
    ],
    "tone": "Blunt",
    "interrupt_tendency": 1.4,
    "follow_up_depth": 7,
    "aggression": 0.6,
    "signature_move": "Cohort by cohort",
    "method": "Opens on retention. I recommend you raise at $12M. Then burn.",
    "question_seeds": [
        "Which cohort expanded most?",
        "Is a valuation of $20M fair?",
        "Which cohort expanded most?",
    ],
    "dealbreakers": ["Services revenue dressed up as recurring"],
}


async def test_builder_drafts_a_persona_with_gemini(live_llm):
    rec = live_llm.install(by_task(persona=[json_reply(DRAFT)]))
    p, source = await builder.generate("A blunt SaaS partner obsessed with NRR who hates services revenue")
    assert source == "gemini" and rec.calls[0]["body"]["model"] == "gemini-3.8-flash"
    assert p.id == "draft" and p.name == "Rhea" and p.archetype == "nrr_obsessed_saas_partner"
    assert p.priorities["retention"] == 1.0 and p.priorities["team"] == 0.0 and p.priorities["market"] == 0.2
    assert p.style.interrupt_tendency == 1.0 and p.style.follow_up_depth == 3 and p.style.tone == "blunt"
    assert p.question_seeds == ["Which cohort expanded most?"]  # money talk and duplicates dropped
    assert p.method == "Opens on retention. Then burn."  # the advice sentence is gone
    assert p.glyph == "custom" and p.kits == []


async def test_builder_falls_back_to_keywords_on_failure(live_llm):
    live_llm.install(by_task(persona=[json_reply({"name": "only a name"})]))  # fails validation
    p, source = await builder.generate(
        "A tough partner named Rhea, obsessed with NRR and churn, hates high burn"
    )
    assert source == "heuristic" and p.name == "Rhea"


async def test_builder_without_a_key_uses_keywords():
    p, source = await builder.generate("A tough partner obsessed with NRR and churn who hates high burn")
    assert source == "heuristic" and p.name == "New investor"
    assert p.priorities["retention"] > 0.6 and p.priorities["runway"] > 0.6 and p.priorities["market"] == 0.2
    assert p.style.aggression > 0.4 and p.question_seeds


def test_builder_options_cover_every_lens_and_role():
    o = builder.options()
    assert len(o["lenses"]) == 12 and "custom" in o["glyphs"] and "lender" in o["cheque_roles"]


# ---------------------------------------------------------------- FR-7 document facts

PAGE2 = ["ARR reached $1.2M in December 2025.", "Monthly burn was $85,000.", "We plan 40% growth in 2026."]


def _fact(name, value, unit, cited, kind="historical", page=None, currency=None, period=None):
    return {
        "name": name,
        "label": name.replace("_", " "),
        "value": value,
        "unit": unit,
        "currency": currency,
        "period": period,
        "kind": kind,
        "page": page,
        "cited_text": cited,
    }


async def test_pdf_facts_are_found_on_their_page_verified_and_pending(live_llm):
    pdf = make_pdf([["Board update"], PAGE2])
    facts = [
        _fact("arr", 1200000, "money", PAGE2[0], page=2, currency="CAD", period="2025-12"),
        _fact("burn", 85000, "money", PAGE2[1], page=1),  # the wrong page: code finds it on page 2
        _fact("growth", 40, "pct", PAGE2[2], kind="forecast", period="2026", page=2),
        _fact("customers", 75, "count", PAGE2[0], page=2),  # 75 isn't in the passage: model-supplied, dropped
        _fact("headcount", 12, "count", "We have 12 people.", page=2),  # a passage that isn't in the document
        _fact("deck_arr", 1200000, "money", PAGE2[0], page=2),  # labelled as a deck's claim: never a record
    ]
    rec = live_llm.install(by_task(doc_facts=[json_reply({"facts": facts}, model="gemini-4-argon")]))
    out = await docs.extract_facts(pdf, "board.pdf", "pdf", upload_id="u1", currency="cad")
    got = {f.name: f for f in out["facts"]}
    assert (
        set(got) == {"arr", "burn", "plan.growth", "headcount"} and out["dropped"] == 2 and out["lines"] == 6
    )
    arr = got["arr"]
    assert arr.status == "pending" and arr.unit == "cad" and arr.value == 1200000
    assert arr.source.kind == "pdf" and arr.source.page == 2 and arr.source.verified is True
    assert arr.source.ref == "board.pdf, p. 2" and arr.source.cited_text == PAGE2[0]
    assert got["burn"].source.page == 2 and got["burn"].source.verified is True  # the page is code's
    assert got["headcount"].source.verified is False and "not found" in got["headcount"].note
    assert got["plan.growth"].value == pytest.approx(0.4) and got["plan.growth"].kind == "forecast"
    (call,) = rec.calls  # a small PDF goes inline, in one call: no Files API
    doc = call["body"]["messages"][0]["content"][0]
    assert doc["type"] == "document" and doc["mime_type"] == "application/pdf" and doc["data"] == pdf
    assert call["body"]["model"] == "gemini-3.8-flash" and call["body"]["response_schema"]["type"] == "object"


async def test_docx_facts_get_their_char_span_from_code(live_llm):
    import docx

    d = docx.Document()
    d.add_paragraph("Quarterly letter")
    d.add_paragraph("Cash in the bank was $483,000 at the end of September 2026.")
    buf = io.BytesIO()
    d.save(buf)
    data = buf.getvalue()
    cited = "Cash in the bank was $483,000 at the end of September 2026."
    rec = live_llm.install(
        by_task(
            doc_facts=[
                json_reply({"facts": [_fact("cash_balance", 483000, "money", cited, period="2026-09")]})
            ]
        )
    )
    out = await docs.extract_facts(data, "letter.docx", "docx", upload_id="u2", currency="usd")
    (f,) = out["facts"]
    assert f.unit == "usd" and f.source.kind == "docx" and f.source.verified is True and f.source.page is None
    content = rec.calls[0]["body"]["messages"][0]["content"]
    assert isinstance(content, str) and cited in content  # text documents go in as text


async def test_a_long_document_is_read_in_parts_and_an_overflowing_part_is_halved(live_llm, monkeypatch):
    monkeypatch.setattr(docs, "PDF_PART_PAGES", 2)
    pdf = make_pdf([[f"Page {i} has {i} customers."] for i in range(1, 6)])
    seen = []

    def respond(call):
        doc = call["body"]["messages"][0]["content"][0]
        from pypdf import PdfReader

        n = len(PdfReader(io.BytesIO(doc["data"])).pages)
        first = int(PdfReader(io.BytesIO(doc["data"])).pages[0].extract_text().split()[1])
        seen.append((first, n))
        if n == 2 and first == 3:  # pages 3-4 overflow once, then are read a page at a time
            return message([{"text": '{"facts": ['}], stop_reason="max_tokens")
        return json_reply(
            {"facts": [_fact("customers", first, "count", f"Page {first} has {first} customers.", page=1)]}
        )

    live_llm.install(respond)
    out = await docs.extract_facts(pdf, "pack.pdf", "pdf", upload_id="u3", currency="usd")
    assert sorted(seen) == [(1, 2), (3, 1), (3, 2), (4, 1), (5, 1)]
    pages = sorted(f.source.page for f in out["facts"])
    assert pages == [1, 3, 4, 5] and all(f.source.verified for f in out["facts"])


def test_text_parts_and_merge_pending():
    text = "\n\n".join(f"Paragraph {i}. " + "x" * 50 for i in range(10))
    parts = docs.text_parts(text, size=200)
    assert "".join(parts) == text and all(len(p) <= 200 for p in parts)
    f = docs.to_fact(
        DocFact(**_fact("arr", 1, "money", "ARR 1")),
        docs.CitedSpan("ARR 1", 1, 2),
        True,
        "",
        upload_id="u",
        filename="x.pdf",
        kind="pdf",
        currency="usd",
    )
    merged, added = docs.merge_pending([], [f])
    again, added2 = docs.merge_pending(merged, [f])
    assert added == 1 and added2 == 0 and again == merged
    assert DocFactList(facts=[]).facts == []
