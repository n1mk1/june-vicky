"""Knowledge base edge cases: converter corners (encodings, delimiters, odd JSON, scripts in HTML),
the OCR plumbing with tesseract mocked, zip bombs and path traversal, the number guard's corner
cases, and API behaviour under limits, conflicts, deletes, concurrency and export with originals."""

import asyncio
import io
import zipfile

import pytest

from tests.llm_fakes import make_pdf
from vik.config import get_settings
from vik.kb import convert, ocr
from vik.kb.organize import excerpt, heuristic_related, scrub_numbers
from vik.kb.store import Resolver, clean_path, clean_tags

# ---------------------------------------------------------------- converters


def test_csv_encodings_and_delimiters():
    (d,) = convert.convert("eu.csv", "Monat;Umsatz\n2025-01;1.000\n".encode("cp1252"))
    assert d.content.splitlines() == ["Monat,Umsatz", "2025-01,1.000"]
    (u,) = convert.convert("u16.csv", "name,city\nZoë,Zürich\n".encode("utf-16"))
    assert "Zoë,Zürich" in u.content
    (q,) = convert.convert("quoted.csv", b'a,b\n"1,000","say ""hi"""\n')
    assert q.content.splitlines()[1] == '"1,000","say ""hi"""'
    with pytest.raises(convert.Unreadable):
        convert.convert("empty.csv", b"\n\n , \n")


def test_html_drops_scripts_and_keeps_tables():
    (d,) = convert.convert(
        "p.html",
        b"<title>T</title><script>alert(1)</script><style>x{}</style>"
        b"<p>Body text</p><table><tr><td>a|b</td><td>2</td></tr></table>",
    )
    assert "alert" not in d.content and "x{}" not in d.content
    assert "Body text" in d.content and "| a\\|b | 2 |" in d.content


def test_json_that_is_not_records_stays_readable():
    (d,) = convert.convert("cfg.json", b'{"plan": {"tiers": [1, 2]}, "name": "x"}')
    assert d.format == "md" and "```json" in d.content and '"tiers"' in d.content
    (bad,) = convert.convert("bad.json", b"{not json")
    assert bad.format == "md" and "{not json" in bad.content


def test_yaml_and_unknown_text_are_fenced_binary_is_a_stub():
    (y,) = convert.convert("x.yaml", b"a: 1\n")
    assert "```yaml" in y.content
    (t,) = convert.convert("notes.weird", b"plain words")
    assert "plain words" in t.content
    (b,) = convert.convert("x.exe", b"MZ\x00\x00" + bytes(range(256)))
    assert b.meta["stub"] and "EXE" in b.content


def test_xlsx_chart_without_cached_values_names_its_cells():
    import openpyxl
    from openpyxl.chart import BarChart, Reference

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Month", "Revenue"])
    ws.append(["Jan", 10])
    ws.append(["Feb", 12])
    ch = BarChart()
    ch.title = "Revenue"
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    ws.add_chart(ch, "D2")
    b = io.BytesIO()
    wb.save(b)
    names = {d.name: d for d in convert.convert("book.xlsx", b.getvalue())}
    assert set(names) == {"book.csv", "book - charts.md"}
    assert (
        "### Chart: Revenue" in names["book - charts.md"].content
        and "$B$2:$B$3" in names["book - charts.md"].content
    )


def test_pdf_page_without_text_is_marked_when_ocr_is_missing(monkeypatch):
    from PIL import Image

    img = Image.new("RGB", (600, 400), "white")
    b = io.BytesIO()
    img.save(b, "PDF")
    monkeypatch.setattr(ocr, "available", lambda: False)
    (d,) = convert.convert("scan.pdf", b.getvalue())
    assert d.meta["unread_pages"] == [1] and d.meta["ocr_pages"] == []
    assert "OCR isn't installed" in d.content and "_No text found on this page._" in d.content


def test_pdf_ocr_path_with_tesseract_mocked(monkeypatch):
    from PIL import Image

    img = Image.new("RGB", (600, 400), "white")
    b = io.BytesIO()
    img.save(b, "PDF")
    monkeypatch.setattr(ocr, "available", lambda: True)
    monkeypatch.setattr(ocr, "image_text", lambda im, sparse=False: "Term sheet: seed round")
    (d,) = convert.convert("scan.pdf", b.getvalue())
    assert d.meta["ocr_pages"] == [1] and d.meta["unread_pages"] == []
    assert "OCR on page 1" in d.content and "Term sheet: seed round" in d.content


def test_text_pdf_is_not_sent_through_ocr(monkeypatch):
    calls = []
    monkeypatch.setattr(ocr, "available", lambda: True)
    monkeypatch.setattr(ocr, "pdf_pages_as_images", lambda *a, **k: calls.append(a) or iter(()))
    (d,) = convert.convert("memo.pdf", make_pdf([["A text layer with plenty of characters on it."]]))
    assert calls == [] and d.meta["ocr_pages"] == []


def test_image_ocr_prefers_the_fuller_reading(monkeypatch):
    from PIL import Image

    b = io.BytesIO()
    Image.new("RGB", (300, 200), "white").save(b, "PNG")
    monkeypatch.setattr(ocr, "available", lambda: True)
    monkeypatch.setattr(
        ocr, "bytes_text", lambda data, sparse=False: "$42,000\nJan\nFeb" if sparse else "Jan"
    )
    (d,) = convert.convert("chart.png", b.getvalue())
    assert "$42,000" in d.content and d.meta["ocr"] is True


def test_zip_bomb_and_traversal_guards(monkeypatch):
    monkeypatch.setattr(get_settings(), "kb_unzip_max_bytes", 1000)
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("big.txt", "0" * 5000)
    with pytest.raises(convert.Unreadable, match="expands to more than"):
        convert.expand("bomb.zip", b.getvalue())
    with pytest.raises(convert.Unreadable):
        convert.expand("fake.zip", b"not a zip")
    folder, stem = convert.split_name("../../etc/pass:wd.md")
    assert folder == "etc" and stem == "pass wd"
    assert clean_path("../../a/<b>/c", "csv") == "a/b/c.csv"


# ---------------------------------------------------------------- guard, links, tags


@pytest.mark.parametrize(
    ("text", "sources", "kept"),
    [
        ("Revenue is $1,200 a month.", "revenue,1200", True),  # thousands separator vs plain
        ("Margin was 72.5%.", "gross margin 72.5%", True),
        ("Margin was 72.50%.", "gross margin 72.5%", True),  # trailing zero
        ("Margin was 73%.", "gross margin 72.5%", False),
        ("See [[FY2024 plan]] next.", "", True),  # numbers in a link are note names
        ("Up 3x since launch.", "", False),
    ],
)
def test_scrub_numbers_corners(text, sources, kept):
    out, dropped = scrub_numbers(text, sources)
    assert (out == text) is kept and (dropped == 0) is kept


def test_resolver_prefers_exact_paths_and_tags_are_cleaned():
    r = Resolver([{"path": "a/Plan.md", "title": "Plan"}, {"path": "b/Plan.md", "title": "Other plan"}])
    assert r.resolve("b/Plan") == "b/Plan.md" and r.resolve("Plan") == "a/Plan.md"
    assert clean_tags(["#Q3 Plan", "q3-plan", "", "Ünïcode!", "財務/計画", "a b!!c"]) == [
        "q3-plan",
        "ünïcode",
        "財務/計画",
        "a-b-c",
    ]


def test_related_links_tables_with_shared_columns_and_excerpts_skip_figures():
    a = {
        "path": "a.csv",
        "title": "a",
        "format": "csv",
        "content": "month,revenue,cogs\n1,2,3\n",
        "kind": "table",
    }
    b = {
        "path": "b.csv",
        "title": "b",
        "format": "csv",
        "content": "month,revenue,opex\n1,2,3\n",
        "kind": "table",
    }
    c = {
        "path": "c.md",
        "title": "c",
        "format": "md",
        "content": "Unrelated legal agreement text.",
        "kind": "note",
    }
    rel = heuristic_related([a, b, c])
    assert rel["a.csv"] == ["b.csv"] and "c.md" not in rel["b.csv"]
    assert "[figure text]" in excerpt({"format": "md", "content": "x\n```text\nOCR junk\n```\ny"})


# ---------------------------------------------------------------- API

pytestmark_db = pytest.mark.db


async def _login(api_client):
    if not (await api_client.get("/auth/session")).json()["ok"]:
        assert (await api_client.post("/auth/access", json={"code": "vik-demo"})).status_code == 200


async def _import(api_client, files):
    from vik import jobs

    r = await api_client.post("/kb/import", files=files)
    assert r.status_code == 200, r.text
    if r.json()["job_id"]:
        await jobs.wait_for([r.json()["job_id"]], timeout_s=60)
    return r.json()


async def _wipe(api_client):
    for n in (await api_client.get("/kb/notes")).json():
        await api_client.delete(f"/kb/notes/{n['id']}")


@pytestmark_db
async def test_kb_api_limits_errors_and_failures(api_client, monkeypatch):
    await _login(api_client)
    await _wipe(api_client)
    assert (await api_client.post("/kb/autofill")).status_code == 422  # empty vault
    assert (await api_client.get("/kb/notes/not-a-uuid")).status_code == 404
    assert (await api_client.get("/kb/notes/00000000-0000-0000-0000-000000000000")).status_code == 404
    assert (await api_client.get("/kb/files/nope")).status_code == 404
    assert (await api_client.get("/kb/search", params={"q": "  "})).json() == []

    monkeypatch.setattr(get_settings(), "kb_upload_max_bytes", 50)
    r = await api_client.post(
        "/kb/import", files=[("files", ("big.csv", b"a,b\n" + b"1,2\n" * 50, "text/csv"))]
    )
    assert r.status_code == 413
    monkeypatch.undo()

    out = await _import(
        api_client,
        [
            ("files", ("broken.xlsx", b"PK\x03\x04 not really a workbook", "application/octet-stream")),
            ("files", ("blob.bin", bytes(range(256)) * 4, "application/octet-stream")),
            ("files", ("ok.csv", b"a,b\n1,2\n", "text/csv")),
        ],
    )
    assert [f["status"] for f in out["files"]] == ["stored", "stored", "stored"]
    status = (await api_client.get("/kb/status")).json()
    # A damaged file fails; a file of a type no converter reads still enters the vault as a stub note.
    assert [f["filename"] for f in status["failures"]] == ["broken.xlsx"]
    assert status["jobs"][0]["result"]["failed"] == 1 and status["jobs"][0]["result"]["notes"] == 2


@pytestmark_db
async def test_kb_notes_crud_paths_and_deletes(api_client):
    await _login(api_client)
    await _wipe(api_client)
    a = (await api_client.post("/kb/notes", json={"title": "Plan", "path": "company/Plan"})).json()
    b = (await api_client.post("/kb/notes", json={"title": "Plan", "path": "company/Plan"})).json()
    assert a["path"] == "company/Plan.md" and b["path"] == "company/Plan 2.md"  # unique paths
    t = (
        await api_client.post("/kb/notes", json={"title": "Grid", "format": "csv", "content": "x,y\n1,2\n"})
    ).json()
    assert t["kind"] == "table" and t["path"] == "Grid.csv"

    # related links follow a move, and vanish when the target is deleted
    await api_client.put(f"/kb/notes/{a['id']}", json={"links": [t["path"]]})
    moved = (await api_client.put(f"/kb/notes/{t['id']}", json={"path": "data/Grid"})).json()
    assert moved["path"] == "data/Grid.csv"
    assert (await api_client.get(f"/kb/notes/{a['id']}")).json()["links"] == ["data/Grid.csv"]
    await api_client.delete(f"/kb/notes/{t['id']}")
    assert (await api_client.get(f"/kb/notes/{a['id']}")).json()["links"] == []
    assert (await api_client.delete(f"/kb/notes/{t['id']}")).status_code == 404

    tags = (
        await api_client.put(f"/kb/notes/{a['id']}", json={"tags": ["#Board", "board", "Q3 Plan"]})
    ).json()["tags"]
    assert tags == ["board", "q3-plan"]

    # An original stays while any note made from it is left.
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "One"
    wb.active.append(["a", "b"])
    wb.create_sheet("Two").append(["c", "d"])
    buf = io.BytesIO()
    wb.save(buf)
    await _import(api_client, [("files", ("two sheets.xlsx", buf.getvalue(), "application/octet-stream"))])
    sheets = [n for n in (await api_client.get("/kb/notes")).json() if n["title"].startswith("two sheets")]
    assert len(sheets) == 2 and sheets[0]["file_id"] == sheets[1]["file_id"]
    fid = sheets[0]["file_id"]
    await api_client.delete(f"/kb/notes/{sheets[0]['id']}")
    assert (await api_client.get(f"/kb/files/{fid}")).status_code == 200
    await api_client.delete(f"/kb/notes/{sheets[1]['id']}")
    assert (await api_client.get(f"/kb/files/{fid}")).status_code == 404

    # The organizer's hubs forget a deleted note: no ghost of it stays in the graph.
    await _import(api_client, [("files", ("ghost check.csv", b"month,revenue\n2025-01,5\n", "text/csv"))])
    victim = next(n for n in (await api_client.get("/kb/notes")).json() if n["title"] == "ghost check")
    hubs = lambda g: [n for n in g["nodes"] if n["kind"] == "missing" and n["title"] == "ghost check"]  # noqa: E731
    await api_client.delete(f"/kb/notes/{victim['id']}")
    assert hubs((await api_client.get("/kb/graph")).json()) == []


@pytestmark_db
async def test_kb_zip_folders_concurrency_and_export_with_originals(api_client):
    await _login(api_client)
    await _wipe(api_client)
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr("finance/q1/p&l.csv", "month,revenue\n2025-01,10\n")
        z.writestr("../escape.md", "# Escape\n")
        z.writestr("__MACOSX/finance/._p&l.csv", "junk")
    # Two imports at once: the organizer lock keeps hubs single.
    r1, r2 = await asyncio.gather(
        api_client.post("/kb/import", files=[("files", ("dump.zip", b.getvalue(), "application/zip"))]),
        api_client.post(
            "/kb/import", files=[("files", ("solo.md", b"# Solo\nA founder note.", "text/markdown"))]
        ),
    )
    from vik import jobs

    await jobs.wait_for([r1.json()["job_id"], r2.json()["job_id"]], timeout_s=60)
    notes = {n["path"]: n for n in (await api_client.get("/kb/notes")).json()}
    assert "finance/q1/p&l.csv" in notes  # the dump's own folders are kept, not re-filed
    assert not any(p.startswith("..") or "/../" in p for p in notes)
    assert sum(1 for p in notes if p == "Home.md") == 1
    hubs = [n for n in notes.values() if n["kind"] == "hub"]
    assert len({h["path"] for h in hubs}) == len(hubs)

    r = await api_client.get("/kb/export", params={"originals": "true"})
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        names = z.namelist()
    assert any(n.startswith("_originals/") and n.endswith("p&l.csv") for n in names)
    assert "finance/q1/p&l.csv" in names and ".vik/meta.json" in names
