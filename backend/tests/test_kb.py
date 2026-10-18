"""Knowledge base: converters (code reads every number), the number guard on model text, the offline
organizer and profile heuristics, and the REST flow import → organize → graph → edit → export →
re-import → autofill (offline agents, on the throwaway database)."""

import io
import zipfile

import pytest

from tests import llm_fakes
from tests.llm_fakes import json_reply, queue
from vicky.kb import autofill, convert
from vicky.kb.autofill import Evidence, heuristic_profile, number_in, propose_profile, verify
from vicky.kb.organize import claude_plan, guess_folder, heuristic_hubs, scrub_numbers
from vicky.kb.store import Resolver, clean_path, wikilinks
from vicky.samples import SAMPLES_DIR


def make_pdf(lines: list[str]) -> bytes:
    return llm_fakes.make_pdf([lines])


def make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Traction"
    cd = CategoryChartData()
    cd.categories = ["Q1", "Q2", "Q3"]
    cd.add_series("ARR", (120000, 180000, 260000))
    s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(1), Inches(1.5), Inches(6), Inches(4), cd)
    s.notes_slide.notes_text_frame.text = "Raising a seed round."
    b = io.BytesIO()
    prs.save(b)
    return b.getvalue()


# ---------------------------------------------------------------- converters


def test_csv_and_tsv_normalize_to_comma_separated():
    (d,) = convert.convert("numbers.tsv", b"month\trevenue\n2025-01\t1000\n\n2025-02\t1200\n")
    assert d.format == "csv" and d.kind == "table"
    assert d.content == "month,revenue\n2025-01,1000\n2025-02,1200\n"
    assert d.meta["rows"] == 2 and d.meta["columns"] == ["month", "revenue"]


def test_xlsx_sheet_per_note():
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "P&L"
    wb.active.append(["Month", "Revenue"])
    wb.active.append(["2025-01", 1000.0])
    other = wb.create_sheet("Empty")
    other.append([None])
    b = io.BytesIO()
    wb.save(b)
    drafts = convert.convert("fin.xlsx", b.getvalue())
    assert [d.name for d in drafts] == ["fin - P&L.csv"]
    assert drafts[0].content == "Month,Revenue\n2025-01,1000\n"  # 1000.0 written as the integer it is


def test_pdf_text_layer_by_page():
    (d,) = convert.convert("memo.pdf", make_pdf(["Investment memo", "Revenue grew 40% this year."]))
    assert d.kind == "document" and d.meta["pages"] == 1
    assert "## Page 1" in d.content and "Revenue grew 40% this year." in d.content


def test_pptx_chart_values_come_from_the_chart_xml():
    (d,) = convert.convert("deck.pptx", make_pptx())
    assert "## Slide 1: Traction" in d.content
    assert "| Q2 | 180000 |" in d.content  # the cached series values, not a model's reading
    assert "Speaker notes: Raising a seed round." in d.content
    assert d.meta["charts"] == 1


def test_docx_keeps_order_headings_and_tables():
    import docx

    doc = docx.Document()
    doc.add_heading("Board update", 1)
    doc.add_paragraph("We grew this quarter.")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Cash", "$1.2M"
    t.cell(1, 0).text, t.cell(1, 1).text = "Runway", "18 months"
    b = io.BytesIO()
    doc.save(b)
    (d,) = convert.convert("board/update.docx", b.getvalue())
    assert d.folder == "board"
    body = d.content
    assert body.index("## Board update") < body.index("We grew") < body.index("| Cash | $1.2M |")


def test_json_records_become_a_table_and_markdown_keeps_frontmatter():
    (d,) = convert.convert("x.json", b'[{"a": 1, "b": "x"}, {"a": 2, "c": true}]')
    assert d.format == "csv" and d.content.splitlines()[0] == "a,b,c"
    (m,) = convert.convert("n.md", b"---\ntitle: My note\ntags: [plan, q3]\n---\n# Heading\nbody")
    assert m.title == "My note" and m.tags == ["plan", "q3"] and m.content.startswith("# Heading")


def test_zip_expands_with_folders_and_skips_junk():
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr("data/a.csv", "x,y\n1,2\n")
        z.writestr("__MACOSX/data/._a.csv", "junk")
        z.writestr(".DS_Store", "junk")
    assert [n for n, _ in convert.expand("dump.zip", b.getvalue())] == ["data/a.csv"]


def test_binary_without_text_becomes_a_stub_note():
    # Any file enters the vault: one no converter reads is a note naming it, linked to its original.
    (d,) = convert.convert("blob.bin", bytes(range(256)) * 10)
    assert d.meta["stub"] and d.format == "md" and "blob.bin" in d.content


# ---------------------------------------------------------------- the number guard and heuristics


def test_scrub_drops_sentences_with_numbers_not_in_the_sources():
    text = "Revenue grew 40% last year. Churn is 3% a month. See [[FY24 plan]] for detail.\n1. Plan for 2025"
    out, dropped = scrub_numbers(text, "Revenue grew 40% in 2025")
    assert "40%" in out and "3%" not in out and "[[FY24 plan]]" in out and "2025" in out
    assert dropped == 1


def test_wikilinks_resolve_like_obsidian():
    notes = [
        {"path": "financials/P&L 2025.csv", "title": "Income statement"},
        {"path": "Home.md", "title": "Home"},
    ]
    r = Resolver(notes)
    assert wikilinks("see [[P&L 2025]] and [[Income statement|the P&L]] and [[home#top]]") == [
        "P&L 2025",
        "Income statement",
        "home",
    ]
    assert r.resolve("P&L 2025") == r.resolve("Income statement") == "financials/P&L 2025.csv"
    assert r.resolve("financials/p&l 2025") == "financials/P&L 2025.csv" and r.resolve("nope") is None
    assert clean_path("a/../b:c/note", "md") == "a/b c/note.md"


def test_heuristic_folder_and_hubs():
    pl = {
        "path": "inbox/x.csv",
        "title": "Monthly revenue and burn",
        "format": "csv",
        "content": "month,revenue\n",
        "kind": "table",
        "summary": "",
    }
    nda = {
        "path": "inbox/y.md",
        "title": "Mutual NDA",
        "format": "md",
        "content": "This agreement ...",
        "kind": "document",
        "summary": "",
    }
    assert guess_folder(pl) == "financials" and guess_folder(nda) == "legal"
    home, hubs = heuristic_hubs([{**pl, "path": "financials/x.csv"}, {**nda, "path": "legal/y.md"}])
    assert [h["title"] for h in hubs] == ["Financials", "Legal"] and "[[Financials]]" in home


def test_profile_heuristic_and_evidence_check():
    notes = [
        {
            "path": "a.md",
            "title": "Deck",
            "format": "md",
            "kind": "document",
            "content": "Northwind Labs Inc. builds billing software. Our SaaS ARR is growing.\nWe are raising $2.5M seed round.",
        },
        {
            "path": "t.csv",
            "title": "P&L",
            "format": "csv",
            "kind": "table",
            "content": "Month,Revenue\n2025-01,1000\n2025-02,1100\n2025-03,1200\n",
        },
    ]
    p = heuristic_profile(notes)
    assert p["name"][0] == "Northwind Labs Inc"
    assert (
        "currency" not in p and "country" not in p
    )  # the reporting currency is CAD unless the founder changes it
    assert p["stage"][0] == "seed" and p["business_model"][0] == "recurring"
    assert p["round_size"][0] == 2_500_000 and p["granularity"][0] == "month"
    by = {n["path"]: n for n in notes}
    ok = verify("round_size", Evidence(value="$2.5M", note="a.md", quote="raising $2.5M seed round"), by)
    assert ok and ok[0] == 2_500_000
    # A quote that isn't in the note, or a number the quote doesn't show, is rejected.
    assert (
        verify("round_size", Evidence(value="$3M", note="a.md", quote="raising $2.5M seed round"), by) is None
    )
    assert verify("name", Evidence(value="Acme", note="a.md", quote="Acme Corp builds"), by) is None
    assert number_in(2_500_000, "raise 2.5 million") and not number_in(25, "raise 2.5 million")


# ---------------------------------------------------------------- Claude path (recorded responses)


async def test_claude_plan_keeps_known_paths_and_drops_invented_numbers(live_llm):
    batch = [
        {
            "path": "inbox/a.csv",
            "title": "a",
            "kind": "table",
            "format": "csv",
            "content": "month,revenue\n2025-01,1000\n",
            "meta": {},
        },
        {
            "path": "inbox/b.md",
            "title": "b",
            "kind": "document",
            "format": "md",
            "content": "A board memo.",
            "meta": {},
        },
    ]
    reply = {
        "notes": [
            {
                "path": "inbox/a.csv",
                "title": "Monthly revenue",
                "folder": "financials",
                "tags": ["revenue"],
                "summary": "Monthly revenue by month. Revenue was 1000 in January. It grew 35% after that.",
                "related": ["inbox/b.md", "nowhere/x.md", "inbox/a.csv"],
            },
            {
                "path": "made/up.md",
                "title": "Ghost",
                "folder": "other",
                "tags": [],
                "summary": "",
                "related": [],
            },
        ]
    }
    rec = live_llm.install(queue(json_reply(reply)))
    plans, dropped = await claude_plan(batch, batch)
    by = {p["path"]: p for p in plans}
    a = by["inbox/a.csv"]
    assert a["folder"] == "financials" and a["title"] == "Monthly revenue"
    assert "1000" in a["summary"] and "35%" not in a["summary"] and dropped == 1  # 35 isn't in the table
    assert a["related"] == ["inbox/b.md"]  # unknown paths and itself are dropped
    assert "made/up.md" not in by and "inbox/b.md" in by  # a note the model skipped gets the heuristic plan
    assert rec.calls[0]["body"]["output_config"]["effort"] == "medium"


async def test_claude_profile_values_must_quote_their_note(live_llm):
    notes = [
        {
            "path": "deck.md",
            "title": "Deck",
            "kind": "document",
            "format": "md",
            "content": "Northwind Labs Inc. builds billing tools.\nWe are raising $2.5M on a SAFE.\nBased in Toronto.",
        }
    ]
    ev = lambda v, q, note="deck.md": {"value": v, "note": note, "quote": q}  # noqa: E731
    reply = {
        "name": ev("Northwind Labs Inc.", "Northwind Labs Inc. builds"),
        "one_liner": ev("Billing tools for 4000 shops.", "builds billing tools"),
        "country": ev("CA", "Based in Toronto"),
        "city": ev("Toronto", "Based in Toronto"),
        "stage": ev("seed", "a seed round"),  # not in the note: rejected
        "round_type": ev("SAFE", "on a SAFE"),
        "round_size": ev("$3M", "raising $2.5M"),  # wrong number: rejected
        "business_model": None,
        "granularity": None,
        "fiscal_year_start_month": None,
    }
    live_llm.install(queue(json_reply(reply)))
    found, by = await propose_profile(notes)
    assert by == "claude"
    values = {k: v for k, (v, _) in found.items()}
    assert (
        values["name"] == "Northwind Labs Inc." and values["country"] == "CA" and values["city"] == "Toronto"
    )
    assert values["round_type"] == "safe" and "currency" not in values
    assert (
        values["round_size"] == 2_500_000
    )  # the heuristic's verified value stands; Claude's $3M was rejected
    assert "stage" not in values and "one_liner" not in values  # unquoted, and an invented figure
    assert found["name"][1] == "From [[deck.md]]: “Northwind Labs Inc. builds”"


def test_profile_batches_carry_every_note_in_full():
    long_doc = {
        "path": "deck.md",
        "kind": "document",
        "format": "md",
        "content": "intro " * 100 + "Based in Toronto.",
    }
    table = {
        "path": "pl.csv",
        "kind": "table",
        "format": "csv",
        "content": "month,revenue\n" + "".join(f"2025-{m:02d},{m}000\n" for m in range(1, 13)),
    }
    batches = autofill.note_batches([table, long_doc], budget=300)
    joined = "".join(batches)
    assert all(len(b) <= 300 for b in batches)
    assert "part 1 of" in batches[0] and "Based in Toronto." in joined  # the long note is split, not cut
    assert "2025-12,12000" in joined  # the whole table, not just its header
    assert joined.index("deck.md") < joined.index("pl.csv")  # documents before tables


async def test_claude_profile_reads_every_batch_and_takes_the_majority(live_llm, monkeypatch):
    monkeypatch.setattr(autofill, "BATCH_CHARS", 12_000)
    pad = "Notes. " * 1500  # each note fills its own batch, well past the old 6,000-character cut
    notes = [
        {"path": f"{k}.md", "title": k, "kind": "document", "format": "md", "content": pad + body}
        for k, body in (
            ("a", "Acme Labs Inc. builds tools."),
            ("b", "Acme Labs Inc. again."),
            ("c", "Zenith Corp is a partner. Our office is in Lyon."),
        )
    ]
    ev = lambda v, q, note: {"value": v, "note": note, "quote": q}  # noqa: E731
    empty = dict.fromkeys(autofill.FIELDS)
    replies = {
        "a.md": {**empty, "name": ev("Acme Labs Inc.", "Acme Labs Inc. builds", "a.md")},
        "b.md": {**empty, "name": ev("Acme Labs Inc.", "Acme Labs Inc. again", "b.md")},
        "c.md": {
            **empty,
            "name": ev("Zenith Corp", "Zenith Corp is a partner", "c.md"),
            "city": ev("Lyon", "Our office is in Lyon", "c.md"),
        },
    }

    def respond(call):
        text = call["body"]["messages"][0]["content"]
        text = text if isinstance(text, str) else " ".join(b.get("text", "") for b in text)
        return json_reply(next(r for path, r in replies.items() if f"### {path}" in text))

    rec = live_llm.install(respond)
    found, by = await propose_profile(notes)
    assert by == "claude" and len(rec.calls) == 3
    assert found["name"][0] == "Acme Labs Inc."  # two batches agree over one
    assert found["city"][0] == "Lyon"  # found past the old cut, in the last batch


# ---------------------------------------------------------------- REST flow

pytestmark_db = pytest.mark.db


async def _wait(api_client, job_id):
    from vicky import jobs

    await jobs.wait_for([job_id], timeout_s=60)
    r = await api_client.get(f"/jobs/{job_id}")
    assert r.json()["status"] == "done", r.text
    return r.json()["result"]


@pytestmark_db
async def test_kb_import_organize_graph_edit_export_autofill(api_client):
    if not (await api_client.get("/auth/session")).json()["ok"]:  # run alone: no cookie from test_api yet
        assert (await api_client.post("/auth/access", json={"code": "vicky-demo"})).status_code == 200
    # Files posted to /uploads by earlier tests are mirrored into the vault: start from an empty one.
    from vicky import db

    await db.execute("delete from kb_notes")
    await db.execute("delete from kb_files")
    assert (await api_client.get("/kb/status")).json()["notes"] == 0

    csv_bytes = (SAMPLES_DIR / "harbor-labs" / "financials_quarterly.csv").read_bytes()
    files = [
        ("files", ("financials_quarterly.csv", csv_bytes, "text/csv")),
        (
            "files",
            (
                "memo.pdf",
                make_pdf(["Investment memo", "We raise a seed round on SaaS revenue."]),
                "application/pdf",
            ),
        ),
        ("files", ("deck.pptx", make_pptx(), "application/octet-stream")),
    ]
    r = await api_client.post("/kb/import", files=files)
    assert r.status_code == 200, r.text
    assert [f["status"] for f in r.json()["files"]] == ["stored"] * 3
    result = await _wait(api_client, r.json()["job_id"])
    assert result["notes"] == 3 and result["failed"] == 0 and result["organize"]["by"] == "heuristic"

    notes = (await api_client.get("/kb/notes")).json()
    paths = {n["path"] for n in notes}
    assert not [p for p in paths if p.startswith("inbox/")]  # the organizer filed everything
    assert "Home.md" in paths and any(n["kind"] == "hub" for n in notes)
    table = next(n for n in notes if n["format"] == "csv")
    assert table["path"].startswith("financials/") and table["summary"].startswith("Table with")

    g = (await api_client.get("/kb/graph")).json()
    ids = {n["id"] for n in g["nodes"]}
    assert all(e["source"] in ids and e["target"] in ids for e in g["links"])
    home = next(n for n in g["nodes"] if n["path"] == "Home.md")
    assert home["degree"] >= 1

    # The same bytes again add nothing.
    r = await api_client.post("/kb/import", files=[("files", ("copy.csv", csv_bytes, "text/csv"))])
    assert r.json()["files"][0]["status"] == "duplicate" and r.json()["job_id"] is None

    # A founder note linking the table, then an edit with a stale version.
    created = (
        await api_client.post(
            "/kb/notes", json={"title": "Plan", "content": f"See [[{table['title']}]] and [[Ghost]]"}
        )
    ).json()
    assert created["author"] == "founder" and created["outlinks"][0]["id"] == table["id"]
    assert any(
        b["id"] == created["id"]
        for b in (await api_client.get(f"/kb/notes/{table['id']}")).json()["backlinks"]
    )
    g = (await api_client.get("/kb/graph")).json()
    assert any(n["kind"] == "missing" and n["title"] == "Ghost" for n in g["nodes"])
    upd = await api_client.put(
        f"/kb/notes/{created['id']}", json={"content": "v2", "version": created["version"]}
    )
    assert upd.status_code == 200 and upd.json()["version"] == created["version"] + 1
    stale = await api_client.put(
        f"/kb/notes/{created['id']}", json={"content": "v3", "version": created["version"]}
    )
    assert stale.status_code == 409
    moved = (await api_client.put(f"/kb/notes/{created['id']}", json={"path": "company/Plan"})).json()
    assert moved["path"] == "company/Plan.md"

    hits = (await api_client.get("/kb/search", params={"q": "Investment memo"})).json()
    assert hits and hits[0]["kind"] == "document"

    # Export → wipe → re-import restores paths, titles and tags.
    r = await api_client.get("/kb/export")
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        names = set(z.namelist())
        assert ".vicky/meta.json" in names and table["path"] in names
        assert z.read("company/Plan.md").decode().startswith("---\ntitle: Plan")
    before = {n["path"]: n for n in (await api_client.get("/kb/notes")).json()}
    for n in before.values():
        assert (await api_client.delete(f"/kb/notes/{n['id']}")).status_code == 200
    assert (await api_client.get("/kb/notes")).json() == []
    r = await api_client.post("/kb/import", files=[("files", ("vault.zip", r.content, "application/zip"))])
    await _wait(api_client, r.json()["job_id"])
    after = {n["path"]: n for n in (await api_client.get("/kb/notes")).json()}
    # Hubs are rewritten after an import (Plan now sits in company/, so a Company hub appears).
    assert set(after) >= set(before) and {p for p in after if after[p]["kind"] != "hub"} == {
        p for p in before if before[p]["kind"] != "hub"
    }
    assert after["company/Plan.md"]["author"] == "founder" and after[table["path"]]["title"] == table["title"]

    # Autofill: a draft engagement with the table mapped and the profile filled, each value cited.
    r = await api_client.post("/kb/autofill")
    assert r.status_code == 200, r.text
    out = r.json()
    eng = out["engagement"]
    assert eng["status"] == "draft" and eng["datasets"] and eng["profile"]["granularity"] == "quarter"
    assert eng["profile_sources"]["granularity"].startswith("From [[financials/")
    assert {a["as"] for a in out["added"]} >= {"financials_quarterly.csv"}
    ds = eng["datasets"][0]
    assert ds["mapped"] >= 0.8 * ds["total"]
